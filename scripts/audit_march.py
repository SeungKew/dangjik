"""Audit March against source workbooks without changing production data.

Convert docs' February/March XLS files to temporary XLSX files first, then run:
  .venv/bin/python scripts/audit_march.py --feb FEB.xlsx --march MARCH.xlsx
Results are written under audit_results/2026-03 by default.
"""

import argparse
import contextlib
import csv
import datetime as dt
import hashlib
import io
import json
from pathlib import Path
import sqlite3
import sys
import tempfile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import dangjik_basic1 as dy
from app import app

TRACKS = {'sen': (1, dy.GROUP_SEN), 'jun': (4, dy.GROUP_JUN),
          'fwd': (7, dy.GROUP_SUK), 'rev': (10, dy.GROUP_SUK)}
# User-confirmed grade; the source row's senior slot is not this person's grade.
PENALTY_EMPLOYEE_GROUPS = {'96100230': dy.GROUP_FIRST}


def read_order(path: Path, month: int) -> list[dict]:
    """Read B:G only; old receipt sheets and helper columns are not orders."""
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    day, kind, slot = 0, '', 0
    for row in wb['당직명령'].iter_rows(min_row=4):
        if isinstance(row[1].value, (int, float)):
            day, slot = int(row[1].value), 0
        if row[3].value in ('일직', '숙직'):
            kind = row[3].value
        if row[5].value is None or row[6].value is None:
            continue
        track = ('sen' if slot == 0 else 'jun') if kind == '일직' else ('rev' if slot else 'fwd')
        rows.append(dict(date=f'2026-{month:02d}-{day:02d}', kind=kind, track=track,
                         name=row[5].value, emp_no=str(row[6].value).zfill(8),
                         penalty=bool(row[5].font.bold)))
        slot += 1
    wb.close()
    return rows


def reference_plan(feb: list[dict], actual: list[dict], path: Path,
                   gap_for_day_duty: bool = False) -> list[dict]:
    """Diagnostic model: March lists, pink exclusions, February final positions.

    The only forced March assignments are the workbook's explicit bold penalties.
    This model is evidence for a repair design, not a production implementation:
    it does not implement later-month skip returns or personnel effective dates.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = wb['순서']
    lists = {}
    for track, (col, _) in TRACKS.items():
        lists[track] = []
        for row in range(3, sheet.max_row + 1):
            cell, number = sheet.cell(row, col), sheet.cell(row, col + 1).value
            if number is None or cell.value is None:
                continue
            lists[track].append(dict(name=cell.value, emp_no=str(number).zfill(8),
                                    excluded=cell.fill.fgColor.type == 'rgb'
                                    and cell.fill.fgColor.rgb == 'FFFF99CC'))
    wb.close()
    pointers = {}
    for track in TRACKS:
        previous = [r for r in feb if r['track'] == track][-1]['emp_no']
        pointers[track] = next(i for i, r in enumerate(lists[track]) if r['emp_no'] == previous)
    history = {r['emp_no']: dt.date.fromisoformat(r['date']) for r in feb}
    forced = {(r['date'], r['track']): r for r in actual if r['penalty']}
    result = []
    kr = dy.holidays.KR(years=[2026])
    for day in range(1, 32):
        date = dt.date(2026, 3, day)
        tracks = ('sen', 'jun', 'rev') if date.weekday() >= 5 or date in kr else ('fwd',)
        for track in tracks:
            key = (date.isoformat(), track)
            if key in forced:
                selected = forced[key]
            else:
                for _ in lists[track]:
                    pointers[track] = (pointers[track] + 1) % len(lists[track])
                    selected = lists[track][pointers[track]]
                    if selected['excluded']:
                        continue
                    last = history.get(selected['emp_no'])
                    if (gap_for_day_duty or track in ('fwd', 'rev')) and last and (date - last).days <= 28:
                        continue
                    break
                else:
                    raise AssertionError(f'No reference candidate: {key}')
            history[selected['emp_no']] = date
            result.append(dict(date=date.isoformat(), kind='일직' if track in ('sen', 'jun') else '숙직',
                               emp_no=selected['emp_no'], name=selected['name']))
    return result


def match_count(actual: list[dict], generated: list[dict]) -> int:
    keys = {(r['date'], r['kind'], r['emp_no']) for r in generated}
    return sum((r['date'], r['kind'], r['emp_no']) in keys for r in actual)


def web_generate() -> list[dict]:
    with contextlib.redirect_stdout(io.StringIO()):
        response = app.test_client().post('/generate', data={'month': '2026-03'}, follow_redirects=True)
    assert response.status_code == 200
    with dy.get_conn() as conn:
        return [dict(date=d, kind=k, emp_no=n, name=name) for d, k, n, name in conn.execute(
            "SELECT h.duty_date,h.duty_type,e.emp_no,e.name FROM duty_history h "
            "JOIN employees e ON h.emp_id=e.id WHERE h.duty_date LIKE '2026-03%' ORDER BY h.id")]


def fingerprint(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--feb', type=Path, required=True)
    parser.add_argument('--march', type=Path, required=True)
    parser.add_argument('--out', type=Path, default=ROOT / 'audit_results/2026-03')
    args = parser.parse_args()
    protected = [Path(dy.V2_DB_NAME), Path(dy.DB_NAME),
                 ROOT / dy.MASTER_FILE, *sorted((ROOT / 'docs').glob('*'))]
    before = {str(p): fingerprint(p) for p in protected if p.is_file()}
    feb, actual = read_order(args.feb, 2), read_order(args.march, 3)
    scenarios = {}
    original_db = dy.DB_NAME
    try:
        with tempfile.TemporaryDirectory(prefix='dangjik-audit-') as tmp:
            dy.DB_NAME = str(Path(tmp) / 'current.db')
            with sqlite3.connect((ROOT / original_db).as_uri() + '?mode=ro', uri=True) as source:
                with sqlite3.connect(dy.DB_NAME) as target:
                    source.backup(target)
            dy.init_db()
            with dy.get_conn() as conn:
                master = {n: dict(name=name, group=g, active=bool(a)) for n, name, g, a in conn.execute(
                    'SELECT emp_no,name,duty_group,is_active FROM employees')}
            scenarios['current_web'] = web_generate()
            scenarios['current_web_repeat'] = web_generate()
            dy.DB_NAME = str(Path(tmp) / 'penalties.db')
            with sqlite3.connect((ROOT / original_db).as_uri() + '?mode=ro', uri=True) as source:
                with sqlite3.connect(dy.DB_NAME) as target:
                    source.backup(target)
            dy.init_db()
            for row in actual:
                if row['penalty'] and row['emp_no'] in master:
                    dy.insert_manual_override(row['date'], PENALTY_EMPLOYEE_GROUPS.get(row['emp_no'], TRACKS[row['track']][1]), row['emp_no'], '벌당직')
            # Direct helper registration is an upper bound: UI hides inactive staff.
            scenarios['known_penalties_only_web'] = web_generate()
            dy.DB_NAME = str(Path(tmp) / 'fresh.db')
            with contextlib.redirect_stdout(io.StringIO()):
                dy.load_initial_data(str(ROOT / dy.MASTER_FILE))
            with dy.get_conn() as conn:
                initial_pointers = conn.execute('SELECT COUNT(*) FROM system_state').fetchone()[0]
            scenarios['fresh_upload_web'] = web_generate()
    finally:
        dy.DB_NAME = original_db
    scenarios['reference'] = reference_plan(feb, actual, args.march)
    scenarios['reference_gap_all_groups'] = reference_plan(feb, actual, args.march, True)
    assert len(feb) == 50 and len(actual) == 51
    assert match_count(actual, scenarios['reference']) == 51
    assert all(fingerprint(Path(p)) == digest for p, digest in before.items()), 'Source asset changed'
    summary = dict(actual_count=len(actual), penalty_count=sum(r['penalty'] for r in actual),
                   missing=[r for r in actual if r['emp_no'] not in master],
                   inactive=[r for r in actual if r['emp_no'] in master and not master[r['emp_no']]['active']],
                   fresh_upload_pointer_count=initial_pointers,
                   generated_counts={k: len(v) for k, v in scenarios.items()},
                   matching_assignments={k: match_count(actual, v) for k, v in scenarios.items()},
                   repeat_same_assignments={(r['date'], r['kind'], r['emp_no']) for r in scenarios['current_web']}
                   == {(r['date'], r['kind'], r['emp_no']) for r in scenarios['current_web_repeat']},
                   production_hashes_unchanged=True, source_sha256=before)
    args.out.mkdir(parents=True, exist_ok=True)
    (args.out / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')
    with (args.out / 'comparison.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['날짜', '분류', '트랙', '실제_이름', '실제_사번', '벌당직', '마스터_상태', *scenarios])
        for r in actual:
            state = master.get(r['emp_no'])
            writer.writerow([r['date'], r['kind'], r['track'], r['name'], r['emp_no'], r['penalty'],
                             '없음' if state is None else '활성' if state['active'] else '제외',
                             *[' / '.join(x['name'] for x in rows if x['date'] == r['date'] and x['kind'] == r['kind'])
                               for rows in scenarios.values()]])
    with (args.out / 'penalties.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['날짜', '당직그룹', '이름', '사번', '사유', '현재_웹_선택가능'])
        for r in actual:
            if r['penalty']:
                writer.writerow([r['date'], PENALTY_EMPLOYEE_GROUPS.get(r['emp_no'], TRACKS[r['track']][1]), r['name'], r['emp_no'], '벌당직',
                                 bool(master.get(r['emp_no'], {}).get('active'))])
    print(json.dumps({k: v for k, v in summary.items() if k not in ('source_sha256', 'inactive')}, ensure_ascii=False, indent=2))
    print(f"Inactive actual assignments: {len(summary['inactive'])}; output: {args.out}")


if __name__ == '__main__':
    main()
