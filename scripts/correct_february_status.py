"""Correct February employee statuses from the source order's pink exclusions.

Supply a temporary XLSX conversion of docs/2026년_2월_본사_당직명령.xls.
Default is a dry run. --apply backs up the template and v3 database first.
Only employee status and template status/reason cells are changed.
"""

import argparse
import csv
import datetime as dt
import json
from pathlib import Path
import shutil
import sqlite3
import tempfile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
GROUP_COLUMNS = {1: '일직_2/3직급_여', 4: '일직_4직급', 7: '숙직_3직급_남'}


def source_statuses(path: Path) -> dict[str, dict]:
    """Pink is exclusion; blue/green/white are not exclusion indicators."""
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = wb['순서']
    result = {}
    try:
        for col, group in GROUP_COLUMNS.items():
            for row in range(3, sheet.max_row + 1):
                cell, number = sheet.cell(row, col), sheet.cell(row, col + 1).value
                if number is None or cell.value is None or str(cell.value).startswith('*'):
                    continue
                number = str(number).strip().zfill(8)
                excluded = cell.fill.patternType == 'solid' and cell.fill.fgColor.type == 'rgb' and cell.fill.fgColor.rgb == 'FFFF99CC'
                info = dict(active=not excluded, group=group, names=[str(cell.value)], cells=[cell.coordinate])
                if number in result:
                    previous = result[number]
                    if previous['active'] != info['active'] or previous['group'] != group:
                        raise ValueError(f'원본의 중복 사번 상태가 다릅니다: {number}')
                    previous['names'].extend(info['names'])
                    previous['cells'].extend(info['cells'])
                else:
                    result[number] = info
        return result
    finally:
        wb.close()


def correct_template(wb: openpyxl.Workbook, source: dict[str, dict]) -> list[dict]:
    sheet = wb['직원명부(마스터)']
    changes = []
    for row in sheet.iter_rows(min_row=2):
        group, _, name, number = [c.value for c in row[:4]]
        if number is None:
            continue
        number = str(number).strip().zfill(8)
        if number == '00000000' or str(name).startswith('*'):
            expected, reason = '제외', '직원 아님(안내행)'
        else:
            info = source[number]
            if group != info['group'] or name not in info['names']:
                raise ValueError(f'입력양식 직원 식별 불일치: {name} / {number}')
            expected = '활성' if info['active'] else '제외'
            reason = None if info['active'] else row[6].value
        if row[5].value != expected or row[6].value != reason:
            changes.append(dict(row=row[0].row, emp_no=number, name=name,
                                old=row[5].value, new=expected, reason=reason))
            row[5].value, row[6].value = expected, reason
    return changes


def database_changes(conn: sqlite3.Connection, source: dict[str, dict]) -> list[dict]:
    changes = []
    for emp_id, number, name, group, active in conn.execute(
            'SELECT id,emp_no,name,duty_group,is_active FROM employees'):
        info = source.get(number)
        if info is None:
            continue  # New arrivals outside the February list are not touched.
        if name not in info['names'] or group != info['group']:
            raise ValueError(f'DB 직원 식별 불일치: {name} / {number}')
        if bool(active) != info['active']:
            changes.append(dict(id=emp_id, emp_no=number, name=name, old=active, new=int(info['active'])))
    return changes


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    source = source_statuses(args.source)
    template = ROOT / '당직명령시스템_2월통합_입력양식.xlsx'
    database = ROOT / 'data/duty_system_v3.db'
    wb = openpyxl.load_workbook(template)
    template_changes = correct_template(wb, source)
    with sqlite3.connect(database.as_uri() + '?mode=ro', uri=True) as conn:
        db_changes = database_changes(conn, source)
    summary = dict(applied=False, source_unique_employees=len(source),
                   source_active=sum(x['active'] for x in source.values()),
                   source_excluded=sum(not x['active'] for x in source.values()),
                   duplicate_source_numbers={n: x for n, x in source.items() if len(x['names']) > 1},
                   template_changes=template_changes, database_changes=db_changes)
    if args.apply and (template_changes or db_changes):
        backup = ROOT / 'backups' / ('employee_status_202602_' + dt.datetime.now().strftime('%Y%m%d_%H%M%S_%f'))
        backup.mkdir(parents=True)
        shutil.copy2(template, backup / template.name)
        with sqlite3.connect(database) as conn:
            with sqlite3.connect(backup / database.name) as copy:
                conn.backup(copy)
            # Keep a full logical snapshot, including all histories and pointers.
            before_other = {r[0]: conn.execute(f'SELECT * FROM "{r[0]}" ORDER BY rowid').fetchall()
                            for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name != 'employees'").fetchall()}
            before_staff = conn.execute('SELECT id,emp_no,name,dept,duty_group,group_idx FROM employees ORDER BY id').fetchall()
            conn.execute('BEGIN IMMEDIATE')
            if database_changes(conn, source) != db_changes:
                raise RuntimeError('검증 후 DB 상태가 변경되었습니다. 다시 실행하세요.')
            for change in db_changes:
                conn.execute('UPDATE employees SET is_active=? WHERE id=?', (change['new'], change['id']))
            assert not database_changes(conn, source)
            assert before_staff == conn.execute('SELECT id,emp_no,name,dept,duty_group,group_idx FROM employees ORDER BY id').fetchall()
            for table, rows in before_other.items():
                assert rows == conn.execute(f'SELECT * FROM "{table}" ORDER BY rowid').fetchall()
            with tempfile.NamedTemporaryFile(dir=template.parent, suffix='.xlsx', delete=False) as staged:
                staged_path = Path(staged.name)
            try:
                wb.save(staged_path)
                verified = openpyxl.load_workbook(staged_path)
                assert not correct_template(verified, source)
                verified.close()
                staged_path.replace(template)
            finally:
                staged_path.unlink(missing_ok=True)
        summary.update(applied=True, backup=str(backup), histories_and_pointers_unchanged=True)
    wb.close()
    out = ROOT / 'audit_results/2026-02'
    out.mkdir(parents=True, exist_ok=True)
    (out / 'status_correction.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')
    with (out / 'status_changes.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['입력양식_행', '성명', '사번', '수정전', '수정후', '제외사유'])
        writer.writerows((r['row'], r['name'], r['emp_no'], r['old'], r['new'], r['reason']) for r in template_changes)
    print(json.dumps({k: v for k, v in summary.items() if k not in ('template_changes', 'database_changes')}, ensure_ascii=False, indent=2))
    print('Template changes:', len(template_changes), 'DB changes:', len(db_changes))


if __name__ == '__main__':
    main()
