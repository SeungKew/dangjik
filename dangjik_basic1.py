"""본사 당직명령 자동 배정 시스템 (연속근무 스킵형 제도).

배정 규칙 (기존 날짜순 방식 및 2026년 7~10월 원본 월별 방식):
- 평일(근무일): 숙직 1명 — 숙직_3직급_남 그룹, group_idx 오름차순(순방향, ptr_sukjig_fwd)
- 주말·공휴일: 일직 2명(일직_2/3직급_여 + 일직_4직급 각 1명, group_idx 오름차순)
  + 숙직 1명 — group_idx 내림차순(역방향, ptr_sukjig_rev)
- 원본 월별 방식: 평일 숙직을 먼저 확정한 뒤 휴일 숙직의 당월 중복과 이전 근무 간격을 판정.
  원본 재현 간격은 30일(관리자 설정 가능). 복귀는 최초 근무·예정 월의 두 달 뒤, 원래 트랙 우선.
- 기존 날짜순 방식: 당월 정책 ON이면 당월, OFF이면 이전 28일 근무를 기준으로 스킵한다.
- 수동 배정(벌당직/대근): manual_overrides 테이블이 최우선.
  벌당직·대근 이력은 이후 정규 배정의 스킵 기준에서 제외(제도 안내 PDF 예외 조항).
"""

import datetime
import os
import shutil
import sqlite3
import calendar
import json
from io import BytesIO
from itertools import groupby
import holidays

import openpyxl

DB_NAME = 'duty_system_v3.db'
V2_DB_NAME = 'duty_system_v2.db'
MASTER_FILE = '당직명령시스템_2월통합_입력양식.xlsx'

# 트랙 → pointer_key 매핑 (duty_system_v2.db의 system_state와 호환)
PTR_SEN = 'ptr_iljig_sen'   # 일직_2/3직급_여 (오름차순)
PTR_JUN = 'ptr_iljig_jun'   # 일직_4직급 (오름차순)
PTR_FWD = 'ptr_sukjig_fwd'  # 숙직 평일 (오름차순)
PTR_REV = 'ptr_sukjig_rev'  # 숙직 주말·공휴일 (내림차순)

GROUP_SEN = '일직_2/3직급_여'
GROUP_JUN = '일직_4직급'
GROUP_SUK = '숙직_3직급_남'
GROUP_FIRST = '일직_1직급'  # 정규 순번 없이 예외 수동 배정에 사용
EMPLOYEE_GROUPS = (GROUP_FIRST, GROUP_SEN, GROUP_JUN, GROUP_SUK)

DUTY_GAP_DAYS = 28          # 당월 스킵 정책 해제 시 사용할 간격
SKIP_RETURN_OFFSET_DAYS = 62  # 스킵자 복귀 시점: 최초 수행(스킵)일 기준 익익월


def get_conn() -> sqlite3.Connection:
    """현재 DB 파일에 대한 커넥션을 반환한다."""
    return sqlite3.connect(DB_NAME)


def init_db() -> None:
    conn = get_conn()
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS employees (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    duty_group TEXT, group_idx INTEGER, name TEXT,
                    emp_no TEXT UNIQUE, dept TEXT, is_active INTEGER DEFAULT 1
                 )''')
    c.execute('''CREATE TABLE IF NOT EXISTS duty_history (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    duty_date DATE, duty_type TEXT, emp_id INTEGER,
                    FOREIGN KEY(emp_id) REFERENCES employees(id)
                 )''')
    c.execute('''CREATE TABLE IF NOT EXISTS skip_queue (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    emp_id INTEGER, skipped_date DATE, return_year_month TEXT,
                    is_returned INTEGER DEFAULT 0, pointer_key TEXT
                 )''')
    _ensure_skip_queue_tracks(c)
    c.execute('''CREATE TABLE IF NOT EXISTS system_state (
                    pointer_key TEXT PRIMARY KEY, emp_id INTEGER
                 )''')
    c.execute('''CREATE TABLE IF NOT EXISTS source_roster_tracks (
                    emp_id INTEGER, pointer_key TEXT, track_idx INTEGER, is_active INTEGER,
                    PRIMARY KEY(emp_id, pointer_key)
                 )''')
    # 수동 배정(벌당직/대근 등) 예외 테이블
    c.execute('''CREATE TABLE IF NOT EXISTS manual_overrides (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    target_date DATE, duty_group TEXT, emp_no TEXT, reason TEXT
                 )''')
    # 월 생성 시작 시점의 포인터 스냅샷 (delete_month 복원용)
    c.execute('''CREATE TABLE IF NOT EXISTS system_state_snapshot (
                    year_month TEXT PRIMARY KEY, snapshot TEXT
                 )''')
    _ensure_employee_change_history(c)
    _ensure_system_settings(c)
    conn.commit()
    conn.close()


def migrate_v2_to_v3() -> None:
    """duty_system_v2.db(기존 정본)를 duty_system_v3.db로 이관한다.

    v3가 없으면 v2 파일을 복사한 뒤 init_db()로 manual_overrides 테이블만 추가한다.
    원본 v2 파일은 데이터 자산이므로 보존한다(수정하지 않는다).
    """
    if os.path.exists(DB_NAME):
        return
    if os.path.exists(V2_DB_NAME):
        shutil.copyfile(V2_DB_NAME, DB_NAME)
        print(f"[이관] {V2_DB_NAME} -> {DB_NAME} 복사 완료")
    init_db()
    print(f"[이관] {DB_NAME} 초기화 완료 (manual_overrides 테이블 포함)")


def reset_database(backup: bool = True) -> str | None:
    """duty_system_v3.db 전체를 삭제하고 빈 스키마로 초기화한다(웹 'DB 초기화' 기능).

    - 백업=True이면 삭제 전 backup_database()로 사본을 생성한다.
    - 빈 DB 파일을 반드시 남겨둔다: 파일만 지우면 재시작 시 migrate_v2_to_v3()가
      v2를 자동 복사해 옛 데이터가 돌아오기 때문이다.
    - v2 원본(duty_system_v2.db)은 데이터 자산이므로 건드리지 않는다.
    반환: 백업 파일 경로(백업 시) 또는 None.
    """
    backup_path = backup_database() if backup else None
    if os.path.exists(DB_NAME):
        os.remove(DB_NAME)
    init_db()
    print(f"[초기화] {DB_NAME} 완전 초기화 완료 (빈 스키마)")
    return backup_path


def backup_database() -> str | None:
    """duty_system_v3.db 사본을 DB와 같은 폴더에 duty_system_v3_backup_<시각>.db 로 생성."""
    if not os.path.exists(DB_NAME):
        return None
    ts = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    backup_path = os.path.join(os.path.dirname(DB_NAME), f"duty_system_v3_backup_{ts}.db")
    shutil.copyfile(DB_NAME, backup_path)
    print(f"[백업] {DB_NAME} -> {backup_path}")
    return backup_path


def load_initial_data(file_path: str) -> None:
    """입력양식 엑셀(직원명부 마스터 + 초기당직이력)을 DB에 동기화한다."""
    init_db()
    wb = openpyxl.load_workbook(file_path)

    # 1) 직원명부(마스터): 사번 기준 upsert
    ws = wb['직원명부(마스터)']
    header = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(header)}
    conn = get_conn()
    c = conn.cursor()
    c.execute('DELETE FROM source_roster_tracks')
    before_employees = {r['id']: r for r in _employee_snapshots(c)}
    added, updated = 0, 0
    removed = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[col['사번']] is None:
            continue
        emp_no = str(row[col['사번']]).strip().zfill(8)
        name = str(row[col['성명']]).strip()
        # 마스터 시트의 주석 행(사번이 0으로만 구성되거나 성명이 '*'로 시작하는
        # 안내/벌당직 예정 표기)은 직원이 아니므로 DB에서 제외한다.
        if set(emp_no.lstrip('0')) == set() or name.startswith('*'):
            c.execute("SELECT id FROM employees WHERE emp_no = ?", (emp_no,))
            existing_noise = c.fetchone()
            if existing_noise is not None:
                emp_id = existing_noise[0]
                depended = c.execute(
                    '''SELECT 1 FROM duty_history WHERE emp_id = ?
                       UNION ALL SELECT 1 FROM skip_queue WHERE emp_id = ?
                       UNION ALL SELECT 1 FROM system_state WHERE emp_id = ?''',
                    (emp_id, emp_id, emp_id)).fetchone()
                if depended is None:
                    c.execute("DELETE FROM employees WHERE id = ?", (emp_id,))
                    removed += 1
                else:
                    c.execute("UPDATE employees SET is_active = 0 WHERE id = ?", (emp_id,))
            continue
        duty_group = row[col['당직그룹']]
        group_idx = int(row[col['해당그룹내순번']])
        dept = row[col['부서']]
        status = str(row[col['상태(활성/제외)']] or '활성').strip()
        is_active = 1 if status == '활성' else 0
        c.execute("SELECT id, name, duty_group, group_idx, is_active FROM employees WHERE emp_no = ?", (emp_no,))
        existing = c.fetchone()
        if existing is None:
            c.execute(
                "INSERT INTO employees (duty_group, group_idx, name, emp_no, dept, is_active) VALUES (?, ?, ?, ?, ?, ?)",
                (duty_group, group_idx, name, emp_no, dept, is_active))
            added += 1
        else:
            c.execute(
                '''UPDATE employees SET duty_group = ?, group_idx = ?, name = ?, dept = ?, is_active = ?
                   WHERE emp_no = ?''',
                (duty_group, group_idx, name, dept, is_active, emp_no))
            updated += 1
    try:
        for after in _employee_snapshots(c):
            before = before_employees.get(after['id'])
            if before != after and not after['name'].startswith('*'):
                _record_employee_change(c, before, after, '초기등록' if before is None else '엑셀변경',
                                        f'입력양식 업로드: {os.path.basename(file_path)}',
                                        employee_change_date(), '엑셀 업로드')
        conn.commit()
    except Exception:
        conn.rollback()
        conn.close()
        wb.close()
        raise

    # 2) 초기당직이력: (당직수행일, 당직분류, 사번) 중복 없이 적재
    ws = wb['초기당직이력(2월)']
    header = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(header)}
    h_added = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[col['당직수행일']] is None:
            continue
        duty_date = str(row[col['당직수행일']]).strip()[:10]
        duty_type = str(row[col['당직분류']]).strip()
        emp_no = str(row[col['사번']]).strip().zfill(8)
        c.execute("SELECT id FROM employees WHERE emp_no = ?", (emp_no,))
        emp = c.fetchone()
        if emp is None:
            print(f"[경고] 마스터에 없는 이력 사번: {emp_no} ({duty_date} {duty_type})")
            continue
        c.execute(
            "SELECT 1 FROM duty_history WHERE duty_date = ? AND duty_type = ? AND emp_id = ?",
            (duty_date, duty_type, emp[0]))
        if c.fetchone() is None:
            c.execute(
                "INSERT INTO duty_history (duty_date, duty_type, emp_id) VALUES (?, ?, ?)",
                (duty_date, duty_type, emp[0]))
            h_added += 1
    conn.commit()
    conn.close()
    wb.close()
    print(f"[불러오기] 직원 신규 {added}명 / 갱신 {updated}명 / 주석행 제거 {removed}명, 이력 신규 {h_added}건")


def import_uploaded_workbook(file_path: str, roster_only: bool = False) -> dict:
    """원본 당직명령과 기존 입력양식을 시트명으로 구분한다."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheets = set(wb.sheetnames)
    finally:
        wb.close()
    if {'당직명령', '순서'} <= sheets:
        return {'format': 'original_roster' if roster_only else 'original',
                **load_original_order(file_path, roster_only=roster_only)}
    if {'직원명부(마스터)', '초기당직이력(2월)'} <= sheets:
        if roster_only:
            raise ValueError('명부만 갱신은 원본 당직명령 파일에서 사용할 수 있습니다.')
        load_initial_data(file_path)
        return {'format': 'template', 'warnings': []}
    raise ValueError('원본 당직명령(당직명령·순서) 또는 기존 입력양식 시트를 확인하세요.')


def _source_emp_no(value: object) -> str:
    text = str(int(value)) if isinstance(value, (int, float)) else str(value or '').strip()
    if not text.isascii() or not text.isdigit() or not 1 <= len(text) <= 8 or int(text) == 0:
        raise ValueError(f'올바르지 않은 사번: {value}')
    return text.zfill(8)


def _source_excluded(cell: openpyxl.cell.cell.Cell) -> bool:
    color = cell.fill.fgColor
    return cell.fill.patternType == 'solid' and (
        (color.type == 'rgb' and str(color.rgb)[-6:].upper() == 'FF99CC')
        or (color.type == 'theme' and color.theme == 5
            and abs(color.tint - 0.7999816888943144) < 0.0001))


def read_original_order(file_path: str) -> dict:
    """원본 명부/명령/스킵을 전부 검증한다. DB에는 쓰지 않는다."""
    import re
    wb = openpyxl.load_workbook(file_path, data_only=True)
    try:
        if not {'순서', '당직명령'} <= set(wb.sheetnames):
            raise ValueError("원본 파일에 '순서', '당직명령' 시트가 필요합니다.")
        order, roster = wb['당직명령'], wb['순서']
        match = re.search(r'(20\d{2})\s*년\s*(\d{1,2})\s*월', ' '.join(str(cell.value or '') for cell in order[1]))
        if not match:
            raise ValueError('당직명령 제목에서 기준 연월을 읽을 수 없습니다.')
        year, month = map(int, match.groups())
        first = datetime.date(year, month, 1)
        ym = first.strftime('%Y-%m')
        employees: dict[str, dict] = {}
        tracks: list[dict] = []
        warnings: list[str] = []
        for col, group, pointer in ((1, GROUP_SEN, PTR_SEN), (4, GROUP_JUN, PTR_JUN),
                                     (7, GROUP_SUK, PTR_FWD), (10, GROUP_SUK, PTR_REV)):
            seen: set[str] = set()
            for row in range(3, roster.max_row + 1):
                cell, raw_no = roster.cell(row, col), roster.cell(row, col + 1).value
                if cell.value is None or str(cell.value).startswith('*') or raw_no is None:
                    continue
                number = _source_emp_no(raw_no)
                name = str(cell.value).strip()
                if number in seen:
                    raise ValueError(f'순서 {cell.coordinate}: 중복 사번 {number}')
                seen.add(number)
                employee = dict(emp_no=number, name=name, dept=str(roster.cell(row, col + 2).value or '').strip(),
                                duty_group=group, group_idx=len(seen), is_active=int(not _source_excluded(cell)))
                if not employee['dept']:
                    raise ValueError(f'순서 {cell.coordinate}: 부서가 없습니다.')
                if pointer != PTR_REV:
                    if number in employees:
                        raise ValueError(f'순서 {cell.coordinate}: 다른 직원과 사번 충돌 {number}')
                    employees[number] = employee
                else:
                    primary = employees.get(number)
                    if primary is None or primary['name'] != name or primary['duty_group'] != group:
                        raise ValueError(f'순서 {cell.coordinate}: 평일/휴일 숙직 직원 식별 불일치 {number}')
                    if primary['is_active'] != employee['is_active']:
                        warnings.append(f'{name}({number}): 평일·휴일 제외 표시가 달라 각각 반영했습니다.')
                # 역순 명부는 위에서 아래가 이미 배정 순서이다.
                tracks.append(dict(emp_no=number, pointer_key=pointer, track_idx=len(seen),
                                   is_active=employee['is_active']))
            if not seen:
                raise ValueError(f'순서 시트의 {group} 명부가 비어 있습니다.')
        if {x['emp_no'] for x in tracks if x['pointer_key'] == PTR_REV} != {
                n for n, e in employees.items() if e['duty_group'] == GROUP_SUK}:
            raise ValueError('평일/휴일 숙직 명부의 직원 집합이 다릅니다.')

        kr = holidays.KR(years=[year])
        duties: list[dict] = []
        date: datetime.date | None = None
        kind = ''
        slots: set[tuple[str, str]] = set()
        for row in range(4, order.max_row + 1):
            raw_day = order.cell(row, 2).value
            if isinstance(raw_day, (int, float)):
                date = datetime.date(year, month, int(raw_day))
            if order.cell(row, 4).value in ('일직', '숙직'):
                kind = order.cell(row, 4).value
            name_value, raw_no = order.cell(row, 6).value, order.cell(row, 7).value
            if name_value is None and raw_no is None:
                continue
            if date is None or kind not in ('일직', '숙직') or name_value is None or raw_no is None:
                raise ValueError(f'당직명령 {row}행: 날짜·분류·성명·사번을 확인하세요.')
            number, name = _source_emp_no(raw_no), str(name_value).strip()
            employee = employees.get(number)
            eligible = lambda e: e['duty_group'] == GROUP_SUK if kind == '숙직' else e['duty_group'] in (GROUP_SEN, GROUP_JUN)
            if employee is None or employee['name'] != name or not eligible(employee):
                candidates = [e for e in employees.values() if e['name'] == name and eligible(e)]
                if len(candidates) != 1:
                    raise ValueError(f'당직명령 {row}행: {name}/{number}를 명부에서 확정할 수 없습니다.')
                corrected = candidates[0]
                warnings.append(f'{date} {kind} {name}: 원본 사번 {number} → 같은 이름의 해당 당직그룹 사번 {corrected["emp_no"]}로 반영했습니다.')
                number, employee = corrected['emp_no'], corrected
            holiday = date.weekday() >= 5 or date in kr
            pointer = PTR_REV if holiday else PTR_FWD
            if kind == '일직':
                pointer = PTR_SEN if employee['duty_group'] == GROUP_SEN else PTR_JUN
            slot = (date.isoformat(), pointer)
            if slot in slots:
                raise ValueError(f'{date}: 같은 당직 자리의 중복 배정입니다.')
            slots.add(slot)
            duties.append(dict(date=date.isoformat(), kind=kind, emp_no=number, pointer_key=pointer,
                               penalty=bool(order.cell(row, 6).font.bold)))
        for day in range(1, calendar.monthrange(year, month)[1] + 1):
            date = datetime.date(year, month, day)
            pointers = (PTR_REV, PTR_SEN, PTR_JUN) if date.weekday() >= 5 or date in kr else (PTR_FWD,)
            if {p for d, p in slots if d == date.isoformat()} != set(pointers):
                raise ValueError(f'{date}: 원본 당직명령이 비어 있거나 당직 자리 수가 맞지 않습니다.')

        def parse_month(value: object) -> str:
            if isinstance(value, (datetime.date, datetime.datetime)):
                return value.strftime('%Y-%m')
            value = str(value or '').strip().replace('.', '-').replace('/', '-')
            parsed = datetime.datetime.strptime(value, '%Y-%m')
            return parsed.strftime('%Y-%m')

        skips: list[dict] = []
        for row in range(3, roster.max_row + 1):
            raw_no = roster.cell(row, 14).value
            if raw_no is None:
                continue
            number = _source_emp_no(raw_no)
            if number not in employees or employees[number]['name'] != str(roster.cell(row, 13).value).strip():
                raise ValueError(f'순서 스킵 {row}행: 직원 성명·사번 불일치')
            skipped, returning = parse_month(roster.cell(row, 16).value), parse_month(roster.cell(row, 17).value)
            if skipped > ym or returning <= skipped:
                raise ValueError(f'순서 스킵 {row}행: 스킵월과 복귀월을 확인하세요.')
            direction = str(roster.cell(row, 19).value or '')
            group = employees[number]['duty_group']
            pointer = {GROUP_SEN: PTR_SEN, GROUP_JUN: PTR_JUN}.get(group)
            if pointer is None:
                if '주말' in direction or '휴일' in direction:
                    pointer = PTR_REV
                elif '주중' in direction or '평일' in direction:
                    pointer = PTR_FWD
                elif returning <= ym:
                    pointer = PTR_FWD  # 이미 복귀 완료한 구월 표는 다음 순번에 영향 없음
                else:
                    raise ValueError(f'순서 스킵 {row}행: 복귀 대기자의 주중/주말 순번 구분이 필요합니다.')
            reason_text = str(roster.cell(row, 18).value or '')
            skip_reason = 'same_month' if '연속근무' in reason_text else 'recent_gap' if '4주' in reason_text else 'imported'
            skips.append(dict(emp_no=number, skipped_date=skipped + '-01', return_year_month=returning, skip_reason=skip_reason,
                              pointer_key=pointer, is_returned=int(returning < ym or (returning == ym and any(
                                  duty['emp_no'] == number and duty['pointer_key'] == pointer for duty in duties)))))
        skip_table_present = str(roster.cell(2, 14).value or '').strip() == '사번'
        if not skip_table_present:
            warnings.append('원본에 스킵 내역 표가 없습니다. 명부만 갱신하여 기존 복귀 대기를 유지하세요.')
        return dict(month=ym, employees=list(employees.values()), tracks=tracks,
                    duties=duties, skips=skips, warnings=warnings, skip_table_present=skip_table_present)
    finally:
        wb.close()


def _sync_original_roster(c: sqlite3.Cursor, data: dict, file_path: str) -> tuple[dict, int]:
    """사번을 유지하며 원본 직원·트랙만 동기화한다. 근무 이력은 읽거나 쓰지 않는다."""
    before = {e['emp_no']: e for e in _employee_snapshots(c)}
    incoming = {e['emp_no'] for e in data['employees']}
    # DB에만 있는 직원은 이력을 유지하되 정규 배정에서 제외한다.
    c.execute('UPDATE employees SET is_active=0')
    max_idx: dict[str, int] = {}
    for employee in data['employees']:
        c.execute('''INSERT INTO employees(emp_no,name,dept,duty_group,group_idx,is_active)
            VALUES (:emp_no,:name,:dept,:duty_group,:group_idx,:is_active)
            ON CONFLICT(emp_no) DO UPDATE SET name=excluded.name,dept=excluded.dept,
            duty_group=excluded.duty_group,group_idx=excluded.group_idx,is_active=excluded.is_active''', employee)
        max_idx[employee['duty_group']] = max(max_idx.get(employee['duty_group'], 0), employee['group_idx'])
    # 소스 밖 직원의 순번 충돌을 피하고 기존 직원 간 순서는 보존한다.
    for old in sorted(before.values(), key=lambda e: (e['group_idx'], e['id'])):
        if old['emp_no'] not in incoming:
            group = old['duty_group']
            max_idx[group] = max_idx.get(group, 0) + 1
            c.execute('UPDATE employees SET group_idx=? WHERE id=?', (max_idx[group], old['id']))
    master = {e['emp_no']: e for e in _employee_snapshots(c)}
    changes = 0
    for number, after in master.items():
        old = before.get(number)
        if old != after:
            logged = _record_employee_change(c, old, after, '초기등록' if old is None else '엑셀변경',
                                    f'원본 {data["month"]} 기준 구성: {os.path.basename(file_path)}',
                                    employee_change_date(), '원본 엑셀 업로드')
            changes += int(logged)
    c.execute('DELETE FROM source_roster_tracks')
    for track in data['tracks']:
        c.execute('INSERT INTO source_roster_tracks(emp_id,pointer_key,track_idx,is_active) VALUES (?,?,?,?)',
                     (master[track['emp_no']]['id'], track['pointer_key'], track['track_idx'], track['is_active']))
    return master, changes


def load_original_order(file_path: str, roster_only: bool = False) -> dict:
    """원본 월말 기준 상태 구성. 검증 후 백업·단일 트랜잭션으로 반영한다."""
    data = read_original_order(file_path)
    if not roster_only and not data['skip_table_present']:
        raise ValueError('스킵 내역 표가 없어 이전 월 복귀 대기를 복원할 수 없습니다. 이전 월을 기준으로 구성한 뒤 이 파일은 직원 명부만 갱신으로 등록하세요.')
    init_db()
    conn = get_conn()
    try:
        stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S_%f')
        backup = os.path.join(os.path.dirname(os.path.abspath(DB_NAME)), f'duty_system_v3_before_source_{stamp}.db')
        with sqlite3.connect(backup) as target:
            conn.backup(target)
        with conn:
            conn.execute('BEGIN IMMEDIATE')
            c = conn.cursor()
            master, changes = _sync_original_roster(c, data, file_path)
            removed = 0
            if not roster_only:
                start = data['month'] + '-01'
                removed = conn.execute('DELETE FROM duty_history WHERE duty_date>=?', (start,)).rowcount
                conn.execute('DELETE FROM system_state_snapshot WHERE year_month>=?', (data['month'],))
                conn.execute('DELETE FROM skip_queue')
                for skip in data['skips']:
                    conn.execute('INSERT INTO skip_queue(emp_id,skipped_date,return_year_month,pointer_key,is_returned,skip_reason) VALUES (?,?,?,?,?,?)',
                                 (master[skip['emp_no']]['id'], skip['skipped_date'], skip['return_year_month'], skip['pointer_key'], skip['is_returned'], skip['skip_reason']))
                consumed_returns = {(s['emp_no'], s['pointer_key']) for s in data['skips']
                                    if s['return_year_month'] == data['month']}
                conn.execute('DELETE FROM system_state')
                for duty in sorted(data['duties'], key=lambda d: d['date']):
                    emp = master[duty['emp_no']]
                    conn.execute('INSERT INTO duty_history(duty_date,duty_type,emp_id) VALUES (?,?,?)',
                                 (duty['date'], duty['kind'], emp['id']))
                    if duty['penalty']:
                        reason = '원본 굵은 표시: 벌당직/대근'
                        exists = conn.execute('SELECT 1 FROM manual_overrides WHERE target_date=? AND duty_group=? AND emp_no=?',
                                              (duty['date'], emp['duty_group'], duty['emp_no'])).fetchone()
                        if not exists:
                            conn.execute('INSERT INTO manual_overrides(target_date,duty_group,emp_no,reason) VALUES (?,?,?,?)',
                                         (duty['date'], emp['duty_group'], duty['emp_no'], reason))
                    elif (duty['emp_no'], duty['pointer_key']) not in consumed_returns:
                        conn.execute('REPLACE INTO system_state(pointer_key,emp_id) VALUES (?,?)', (duty['pointer_key'], emp['id']))
        return dict(month=data['month'], employees=len(data['employees']), duties=0 if roster_only else len(data['duties']),
                    skips=0 if roster_only else len(data['skips']), changes=changes, removed_history=removed,
                    backup=backup, warnings=data['warnings'])
    finally:
        conn.close()


def add_manual_override() -> None:
    """관리자가 특정 날짜에 특정 인원을 수동으로 강제 배정 (벌당직 등)"""
    date_str: str = input("수동 지정할 날짜 (YYYY-MM-DD): ")
    group: str = input(f"당직그룹 ({', '.join(EMPLOYEE_GROUPS)}): ")
    emp_no: str = input("대상자 사번 (8자리): ")
    reason: str = input("사유 (예: 벌당직, 대근): ")

    conn = get_conn()
    c = conn.cursor()
    c.execute("INSERT INTO manual_overrides (target_date, duty_group, emp_no, reason) VALUES (?, ?, ?, ?)",
              (date_str, group, str(emp_no).zfill(8), reason))
    conn.commit()
    conn.close()
    print(f"[{date_str}] {group}에 수동 배정되었습니다.")


def _infer_skip_queue_track(duty_group: str, skipped_date: str) -> str | None:
    """구 큐의 스킵일과 그룹으로 원래 순번 트랙을 복원한다."""
    if duty_group == GROUP_SEN:
        return PTR_SEN
    if duty_group == GROUP_JUN:
        return PTR_JUN
    if duty_group == GROUP_SUK:
        day = datetime.date.fromisoformat(skipped_date)
        return PTR_REV if day.weekday() >= 5 or day in holidays.KR(years=[day.year]) else PTR_FWD
    return None


def _ensure_skip_queue_tracks(c: sqlite3.Cursor) -> None:
    columns = {row[1] for row in c.execute('PRAGMA table_info(skip_queue)')}
    if 'pointer_key' not in columns:
        c.execute('ALTER TABLE skip_queue ADD COLUMN pointer_key TEXT')
    for column in ('skip_reason', 'source_duty_date', 'returned_date'):
        if column not in columns:
            c.execute(f'ALTER TABLE skip_queue ADD COLUMN {column} TEXT')
    rows = c.execute('''SELECT q.id, q.skipped_date, e.duty_group FROM skip_queue q
                        JOIN employees e ON e.id=q.emp_id
                        WHERE q.pointer_key IS NULL OR q.pointer_key='' ''').fetchall()
    for queue_id, date, group in rows:
        track = _infer_skip_queue_track(group, date)
        if track is not None:
            c.execute('UPDATE skip_queue SET pointer_key=? WHERE id=?', (track, queue_id))
    c.execute('''CREATE INDEX IF NOT EXISTS idx_skip_queue_track_return
                 ON skip_queue(pointer_key, return_year_month, is_returned, skipped_date, id)''')


def _ensure_system_settings(c: sqlite3.Cursor) -> None:
    c.execute('''CREATE TABLE IF NOT EXISTS system_settings (
                    setting_key TEXT PRIMARY KEY, setting_value TEXT NOT NULL
                 )''')
    c.execute("INSERT OR IGNORE INTO system_settings VALUES ('skip_current_month_only', '1')")


def _skip_current_month_only(c: sqlite3.Cursor) -> bool:
    # 설정 테이블이 없는 구DB도 기본값은 당월 정책이다.
    if not c.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='system_settings'").fetchone():
        return True
    row = c.execute("SELECT setting_value FROM system_settings WHERE setting_key='skip_current_month_only'").fetchone()
    return row is None or row[0] != '0'


def get_skip_current_month_only() -> bool:
    conn = get_conn()
    try:
        return _skip_current_month_only(conn.cursor())
    finally:
        conn.close()


def set_skip_current_month_only(enabled: bool) -> None:
    conn = get_conn()
    try:
        with conn:
            c = conn.cursor()
            _ensure_system_settings(c)
            c.execute("UPDATE system_settings SET setting_value=? WHERE setting_key='skip_current_month_only'",
                      ('1' if enabled else '0',))
    finally:
        conn.close()


def _has_recent_regular_duty(c: sqlite3.Cursor, emp_id: int, duty_date: str) -> bool:
    """설정에 따라 당월 또는 이전28일의 정규근무로 스킵 여부를 판단한다.

    예외근무도 이력에는 남긴다. 같은 날짜의 중복 배정은 예외와 무관하게 막고,
    과거 월 생성 시 미래 이력은 간격 판정에 사용하지 않는다.
    자유 입력 사유 중 명시된 '벌당직'·'대근'만 예외로 인식한다.
    """
    current = datetime.date.fromisoformat(duty_date)
    start = (current.replace(day=1) if _skip_current_month_only(c)
             else current - datetime.timedelta(days=DUTY_GAP_DAYS)).isoformat()
    return c.execute('''SELECT 1 FROM duty_history h
                       JOIN employees e ON e.id = h.emp_id
                       WHERE h.emp_id = ? AND h.duty_date BETWEEN ? AND ?
                         AND (h.duty_date = ? OR NOT EXISTS (
                             SELECT 1 FROM manual_overrides m
                             WHERE m.target_date = h.duty_date AND m.emp_no = e.emp_no
                               AND (INSTR(COALESCE(m.reason, ''), '벌당직') > 0
                                    OR INSTR(COALESCE(m.reason, ''), '대근') > 0)))
                       LIMIT 1''', (emp_id, start, duty_date, duty_date)).fetchone() is not None


def get_next_candidate(c: sqlite3.Cursor, current_date_str: str, group_name: str,
                       pointer_key: str, reverse: bool = False) -> tuple[int | None, str]:
    """수동 배정 -> 큐 복귀 -> 정규 순번 순으로 탐색"""
    # 1. 수동 배정(Override) 내역이 있는지 최우선 확인
    # 1직급의 예외 일직은 정(2/3직급) 자리를 대체하며 정규 포인터는 진행하지 않는다.
    override_groups = (group_name, GROUP_FIRST if group_name == GROUP_SEN else group_name)
    c.execute('''SELECT e.id, e.name FROM manual_overrides m
                 JOIN employees e ON m.emp_no = e.emp_no
                 WHERE m.target_date = ? AND m.duty_group IN (?, ?) ORDER BY m.id''',
              (current_date_str, *override_groups))
    override = c.fetchone()
    if override:
        return override[0], override[1] + "(수동)"

    current_date = datetime.datetime.strptime(current_date_str, "%Y-%m-%d").date()
    current_ym = current_date.strftime("%Y-%m")

    # 2. 원래 순번 트랙의 익익월 복귀자만 확인한다.
    c.execute('''SELECT q.id, q.emp_id, e.name, q.pointer_key, q.skipped_date FROM skip_queue q
                 JOIN employees e ON q.emp_id = e.id
                 LEFT JOIN source_roster_tracks t ON t.emp_id=e.id AND t.pointer_key=?
                 WHERE e.duty_group = ? AND q.return_year_month = ? AND q.is_returned = 0 AND COALESCE(t.is_active,e.is_active) = 1
                   AND (q.pointer_key = ? OR q.pointer_key IS NULL OR q.pointer_key = '')
                 ORDER BY q.skipped_date ASC, q.id ASC''', (pointer_key, group_name, current_ym, pointer_key))
    for skip_return in c.fetchall():
        track = skip_return[3]
        if not track:
            track = _infer_skip_queue_track(group_name, skip_return[4])
            if track is not None:
                c.execute('UPDATE skip_queue SET pointer_key=? WHERE id=?', (track, skip_return[0]))
        if track != pointer_key:
            continue
        if not _has_recent_regular_duty(c, skip_return[1], current_date_str):
            c.execute("UPDATE skip_queue SET is_returned = 1, returned_date=? WHERE id = ?", (current_date_str, skip_return[0]))
            return skip_return[1], skip_return[2] + "(복귀)"

    # 3. 저장된 스킵 정책에 따라 정규 순번 탐색
    c.execute("SELECT emp_id FROM system_state WHERE pointer_key = ?", (pointer_key,))
    res = c.fetchone()
    last_emp_id = res[0] if res else 0

    order_dir = "DESC" if reverse else "ASC"
    # 전출된 마지막 배정자의 위치도 찾아야 순번이 처음으로 되돌아가지 않는다.
    c.execute(f"""SELECT e.id,e.name,COALESCE(t.is_active,e.is_active) FROM employees e
                  LEFT JOIN source_roster_tracks t ON t.emp_id=e.id AND t.pointer_key=?
                  WHERE e.duty_group=?
                  ORDER BY COALESCE(t.track_idx,
                      CASE WHEN ? THEN (SELECT count(*) FROM employees other
                           WHERE other.duty_group=e.duty_group AND other.group_idx>=e.group_idx)
                           ELSE e.group_idx END) ASC,
                           e.group_idx {order_dir}, e.id {order_dir}""", (pointer_key, group_name, reverse))
    emps = c.fetchall()
    if not any(emp[2] for emp in emps):
        return None, "배정불가"

    start_idx = 0
    for idx, emp in enumerate(emps):
        if emp[0] == last_emp_id:
            start_idx = (idx + 1) % len(emps)
            break

    curr_idx = start_idx
    while True:
        emp_id, name, active = emps[curr_idx]
        if not active:
            curr_idx = (curr_idx + 1) % len(emps)
            if curr_idx == start_idx:
                return None, "배정불가"
            continue
        if not _has_recent_regular_duty(c, emp_id, current_date_str):
            c.execute("REPLACE INTO system_state (pointer_key, emp_id) VALUES (?, ?)", (pointer_key, emp_id))
            return emp_id, name

        return_ym = (current_date.replace(day=1) + datetime.timedelta(days=SKIP_RETURN_OFFSET_DAYS)).strftime("%Y-%m")
        c.execute("INSERT INTO skip_queue (emp_id, skipped_date, return_year_month, pointer_key) VALUES (?, ?, ?, ?)",
                  (emp_id, current_date_str, return_ym, pointer_key))

        curr_idx = (curr_idx + 1) % len(emps)
        if curr_idx == start_idx:
            return None, "배정불가"


def assign_day(c: sqlite3.Cursor, d: datetime.date) -> list[tuple[str, str | None, str]]:
    """하루치 당직을 배정한다. 반환: [(당직분류, 사번orNone, 표기)]"""
    is_nonworking = d.weekday() >= 5 or d in holidays.KR(years=[d.year])
    plan: list[tuple[str, str, bool]] = []
    if is_nonworking:
        plan.append(('일직', GROUP_SEN, False))
        plan.append(('일직', GROUP_JUN, False))
        plan.append(('숙직', GROUP_SUK, True))
    else:
        plan.append(('숙직', GROUP_SUK, False))

    assigned: list[tuple[str, str | None, str]] = []
    for duty_type, group, reverse in plan:
        pointer_key = {GROUP_SEN: PTR_SEN, GROUP_JUN: PTR_JUN, GROUP_SUK: PTR_REV if reverse else PTR_FWD}[group]
        emp_id, label = get_next_candidate(c, d.isoformat(), group, pointer_key, reverse=reverse)
        if emp_id is None:
            print(f"[경고] {d} {duty_type}({group}): 배정불가 — 수동 배정 필요")
            assigned.append((duty_type, None, "(미배정)"))
            continue
        # 수동/자동 모두 기록하되, 이후 간격 판정에서는 벌당직·대근 예외를 구분한다.
        c.execute("INSERT INTO duty_history (duty_date, duty_type, emp_id) VALUES (?, ?, ?)",
                  (d.isoformat(), duty_type, emp_id))
        assigned.append((duty_type, _emp_no_of(c, emp_id), label))
    return assigned


def _emp_no_of(c: sqlite3.Cursor, emp_id: int) -> str | None:
    c.execute("SELECT emp_no FROM employees WHERE id = ?", (emp_id,))
    row = c.fetchone()
    return row[0] if row else None


def _add_months(year_month: str, offset: int) -> str:
    year, month = map(int, year_month.split('-'))
    ordinal = year * 12 + month - 1 + offset
    return f'{ordinal // 12:04d}-{ordinal % 12 + 1:02d}'


def _generation_settings(c: sqlite3.Cursor) -> dict:
    rows = dict(c.execute('SELECT setting_key,setting_value FROM system_settings').fetchall())
    mode = rows.get('generation_mode')
    if mode is None:
        mode = 'source_monthly' if c.execute('SELECT 1 FROM source_roster_tracks LIMIT 1').fetchone() else 'legacy'
    return dict(mode=mode, prior_gap_days=int(rows.get('source_prior_gap_days', '30')),
                current_month_only=_skip_current_month_only(c))


def get_generation_settings() -> dict:
    conn = get_conn()
    try:
        return _generation_settings(conn.cursor())
    finally:
        conn.close()


def set_generation_settings(mode: str, prior_gap_days: int, current_month_only: bool) -> None:
    if mode not in ('source_monthly', 'legacy') or not 0 <= prior_gap_days <= 365:
        raise ValueError('배정 방식과 이전 근무 제한 일수(0~365)를 확인하세요.')
    conn = get_conn()
    try:
        with conn:
            _ensure_system_settings(conn.cursor())
            for key, value in (('generation_mode', mode), ('source_prior_gap_days', str(prior_gap_days)),
                               ('skip_current_month_only', '1' if current_month_only else '0')):
                conn.execute('INSERT OR REPLACE INTO system_settings VALUES (?,?)', (key, value))
    finally:
        conn.close()


def _regular_dates(c: sqlite3.Cursor, emp_id: int, first: str, last: str) -> list[str]:
    return [r[0] for r in c.execute('''SELECT h.duty_date FROM duty_history h
        JOIN employees e ON e.id=h.emp_id
        WHERE h.emp_id=? AND h.duty_date BETWEEN ? AND ?
        AND NOT EXISTS (SELECT 1 FROM manual_overrides m
            WHERE m.emp_no=e.emp_no AND m.target_date=h.duty_date
            AND (m.reason LIKE '%벌당직%' OR m.reason LIKE '%대근%'))
        ORDER BY h.duty_date DESC''', (emp_id, first, last)).fetchall()]


def _source_conflict(c: sqlite3.Cursor, emp_id: int, duty_date: str,
                     pointer_key: str, settings: dict) -> tuple[str, str] | None:
    """충돌 사유와 최초 수행/예정일. 같은 달의 미래 평일 배정도 조회한다."""
    day = datetime.date.fromisoformat(duty_date)
    if c.execute('SELECT 1 FROM duty_history WHERE emp_id=? AND duty_date=?', (emp_id, duty_date)).fetchone():
        return 'same_day', duty_date
    first = day.replace(day=1)
    last = day.replace(day=calendar.monthrange(day.year, day.month)[1])
    dates = _regular_dates(c, emp_id, first.isoformat(), last.isoformat())
    if not settings['current_month_only']:
        dates = [date for date in dates if abs((datetime.date.fromisoformat(date) - day).days) <= DUTY_GAP_DAYS]
    if dates:
        return 'same_month', dates[0]
    gap = settings['prior_gap_days'] if pointer_key in (PTR_FWD, PTR_REV) else 0
    if gap:
        dates = _regular_dates(c, emp_id, (day - datetime.timedelta(days=gap)).isoformat(),
                               (first - datetime.timedelta(days=1)).isoformat())
        if dates:
            return 'recent_gap', dates[0]
    return None


def _choose_source_monthly(c: sqlite3.Cursor, date: str, group: str,
                           pointer: str, settings: dict) -> tuple[int, str]:
    # 벌당직/대근 등 수동 자리는 정규/복귀 순번을 진행하지 않는다.
    groups = (group, GROUP_FIRST if group == GROUP_SEN else group)
    manual = c.execute('''SELECT e.id,e.name FROM manual_overrides m JOIN employees e ON e.emp_no=m.emp_no
        WHERE m.target_date=? AND m.duty_group IN (?,?) ORDER BY m.id''', (date, *groups)).fetchall()
    if manual:
        if len(manual) != 1:
            raise ValueError(f'{date} {group}: 같은 자리에 수동 예외가 여러 건입니다.')
        emp_id, name = manual[0]
        if c.execute('SELECT 1 FROM duty_history WHERE emp_id=? AND duty_date=?', (emp_id, date)).fetchone():
            raise ValueError(f'{date}: {name}의 같은 날 중복 수동 배정입니다.')
        return emp_id, name + '(수동)'
    ym = date[:7]
    # 이번 달 복귀 묶음을 우선하고 미소진된 이전 묶음은 이후 달에도 보존한다.
    returns = c.execute('''SELECT q.id,q.emp_id,e.name FROM skip_queue q JOIN employees e ON e.id=q.emp_id
        LEFT JOIN source_roster_tracks t ON t.emp_id=e.id AND t.pointer_key=?
        WHERE q.pointer_key=? AND q.return_year_month<=? AND q.is_returned=0
        AND e.duty_group=? AND COALESCE(t.is_active,e.is_active)=1
        ORDER BY q.return_year_month DESC,q.skipped_date,q.id''', (pointer, pointer, ym, group)).fetchall()
    for queue_id, emp_id, name in returns:
        if _source_conflict(c, emp_id, date, pointer, settings) is None:
            c.execute('UPDATE skip_queue SET is_returned=1,returned_date=? WHERE id=?', (date, queue_id))
            return emp_id, name + '(복귀)'
    reverse = pointer == PTR_REV
    direction = 'DESC' if reverse else 'ASC'
    rows = c.execute(f'''SELECT e.id,e.name,COALESCE(t.is_active,e.is_active) FROM employees e
        LEFT JOIN source_roster_tracks t ON t.emp_id=e.id AND t.pointer_key=? WHERE e.duty_group=?
        ORDER BY COALESCE(t.track_idx,CASE WHEN ? THEN
            (SELECT count(*) FROM employees other WHERE other.duty_group=e.duty_group AND other.group_idx>=e.group_idx)
            ELSE e.group_idx END),e.group_idx {direction},e.id {direction}''', (pointer, group, reverse)).fetchall()
    last = c.execute('SELECT emp_id FROM system_state WHERE pointer_key=?', (pointer,)).fetchone()
    positions = {row[0]: idx for idx, row in enumerate(rows)}
    start = (positions[last[0]] + 1) % len(rows) if last and last[0] in positions else 0
    for offset in range(len(rows)):
        emp_id, name, active = rows[(start + offset) % len(rows)]
        if not active:
            continue
        conflict = _source_conflict(c, emp_id, date, pointer, settings)
        if conflict is None:
            c.execute('REPLACE INTO system_state VALUES (?,?)', (pointer, emp_id))
            return emp_id, name
        reason, anchor = conflict
        # 복귀월은 최초 정규 근무(예정)월의 두 달 뒤이다.
        returning = _add_months(anchor[:7], 2)
        exists = c.execute('''SELECT 1 FROM skip_queue WHERE emp_id=? AND pointer_key=?
            AND skipped_date LIKE ? AND return_year_month=?''', (emp_id, pointer, ym + '%', returning)).fetchone()
        if not exists:
            c.execute('''INSERT INTO skip_queue(emp_id,skipped_date,return_year_month,pointer_key,skip_reason,source_duty_date)
                         VALUES (?,?,?,?,?,?)''', (emp_id, date, returning, pointer, reason, anchor))
    raise ValueError(f'{date} {group}: 조건을 만족하는 배정 대상자가 없습니다.')


def _generate_source_monthly(start: datetime.date, end: datetime.date) -> None:
    """월 단위로 평일 숙직을 먼저 확정하고 휴일 순번 충돌을 처리한다."""
    if start.day != 1 or end.day != calendar.monthrange(end.year, end.month)[1]:
        raise ValueError('원본 월별 방식은 월 전체를 생성하세요. 부분 기간은 기존 날짜순 방식을 사용하세요.')
    conn = get_conn()
    try:
        with conn:
            conn.execute('BEGIN IMMEDIATE')
            c = conn.cursor()
            if c.execute('SELECT 1 FROM duty_history WHERE duty_date>? LIMIT 1', (end.isoformat(),)).fetchone():
                raise ValueError('생성 기간 이후의 당직 데이터가 있습니다. 이후 월을 먼저 삭제한 뒤 순서대로 생성하세요.')
            settings = _generation_settings(c)
            months = sorted({day.strftime('%Y-%m') for day in days_of_month_iter(start, end)})
            for ym in months:
                days = [day for day in days_of_month_iter(start, end) if day.strftime('%Y-%m') == ym]
                snap = c.execute('SELECT snapshot FROM system_state_snapshot WHERE year_month=?', (ym,)).fetchone()
                if snap is not None:
                    c.execute('DELETE FROM system_state')
                    for pair in snap[0].split(';'):
                        if '=' in pair:
                            key, emp_id = pair.split('=', 1)
                            c.execute('INSERT INTO system_state VALUES (?,?)', (key, int(emp_id)))
                else:
                    snapshot = ';'.join(f'{key}={emp_id}' for key, emp_id in c.execute('SELECT * FROM system_state').fetchall())
                    c.execute('INSERT INTO system_state_snapshot VALUES (?,?)', (ym, snapshot))
                c.execute('DELETE FROM duty_history WHERE duty_date BETWEEN ? AND ?', (days[0].isoformat(), days[-1].isoformat()))
                c.execute('DELETE FROM skip_queue WHERE skipped_date LIKE ?', (ym + '%',))
                c.execute('''UPDATE skip_queue SET is_returned=0,returned_date=NULL
                    WHERE returned_date LIKE ? OR (returned_date IS NULL AND return_year_month=?)''', (ym + '%', ym))
                c.execute('DELETE FROM system_state_snapshot WHERE year_month>?', (ym,))
                kr = holidays.KR(years=[days[0].year])
                weekdays = [day for day in days if day.weekday() < 5 and day not in kr]
                holidays_in_month = [day for day in days if day not in weekdays]
                slots = [(day, GROUP_SUK, PTR_FWD) for day in weekdays]
                slots += [(day, group, pointer) for day in holidays_in_month
                          for group, pointer in ((GROUP_SEN, PTR_SEN), (GROUP_JUN, PTR_JUN), (GROUP_SUK, PTR_REV))]
                for day, group, pointer in slots:
                    emp_id, name = _choose_source_monthly(c, day.isoformat(), group, pointer, settings)
                    kind = '숙직' if group == GROUP_SUK else '일직'
                    c.execute('INSERT INTO duty_history(duty_date,duty_type,emp_id) VALUES (?,?,?)', (day.isoformat(), kind, emp_id))
                    print(f'{day} {kind}: {name}')
    finally:
        conn.close()


def generate_duty_plan(start_str: str, end_str: str) -> None:
    """지정 기간(포함)의 당직 배정을 생성하고 duty_history에 기록한다."""
    start = datetime.datetime.strptime(start_str, "%Y-%m-%d").date()
    end = datetime.datetime.strptime(end_str, "%Y-%m-%d").date()
    if end < start:
        print("종료일이 시작일보다 빠릅니다.")
        return

    init_db()
    if get_generation_settings()['mode'] == 'source_monthly':
        _generate_source_monthly(start, end)
        return
    conn = get_conn()
    c = conn.cursor()
    # 기간 재생성 시 중복 방지: 해당 기간의 기존 배정 기록은 덮어쓴다(수동 배정은 유지).
    c.execute("DELETE FROM duty_history WHERE duty_date BETWEEN ? AND ?", (start_str, end_str))
    conn.commit()
    dow_names = "월화수목금토일"
    snap_month = None
    for one in days_of_month_iter(start, end):
        # 월 첫 진입 시 포인터 스냅샷 기록 (이미 있으면 유지 — delete_month 복원용)
        cur_month = one.strftime("%Y-%m")
        if cur_month != snap_month:
            snap_month = cur_month
            c.execute("SELECT 1 FROM system_state_snapshot WHERE year_month = ?", (cur_month,))
            if c.fetchone() is None:
                state = c.execute("SELECT pointer_key, emp_id FROM system_state").fetchall()
                snap = ';'.join(f"{k}={v}" for k, v in state)
                c.execute("INSERT INTO system_state_snapshot (year_month, snapshot) VALUES (?, ?)",
                          (cur_month, snap))
                conn.commit()
        dow = dow_names[one.weekday()]
        kind = "휴무일" if (one.weekday() >= 5 or one in holidays.KR(years=[one.year])) else "평일"
        for duty_type, emp_no, label in assign_day(c, one):
            emp_show = emp_no or "-"
            print(f"{one}({dow}) [{kind}] {duty_type}: {label} (사번 {emp_show})")
        conn.commit()
    conn.close()
    print(f"\n배정 완료: {start} ~ {end}\n")


def days_of_month_iter(start: datetime.date, end: datetime.date):
    기 = start
    while 기 <= end:
        yield 기
        기 = 기 + datetime.timedelta(days=1)


def show_status() -> None:
    """지정 월의 현재 배정 현황을 조회한다."""
    year_month: str = input("조회할 연월 (YYYY-MM): ")
    conn = get_conn()
    c = conn.cursor()
    c.execute('''SELECT h.duty_date, h.duty_type, e.name, e.emp_no
                 FROM duty_history h JOIN employees e ON h.emp_id = e.id
                 WHERE h.duty_date LIKE ? ORDER BY h.duty_date, h.id''', (year_month + '%',))
    rows = c.fetchall()
    conn.close()
    print(f"== {year_month} 배정 현황 ({len(rows)}건)")
    for duty_date, duty_type, name, emp_no in rows:
        print(f"  {duty_date} {duty_type}: {name} ({emp_no})")


def insert_manual_override(target_date: str, duty_group: str, emp_no: str, reason: str) -> bool:
    """수동 배정(벌당직/대근)을 등록한다. 대상 사번이 없으면 False."""
    emp_no = str(emp_no).strip().zfill(8)
    conn = get_conn()
    c = conn.cursor()
    c.execute("SELECT id, name, duty_group FROM employees WHERE emp_no = ?", (emp_no,))
    emp = c.fetchone()
    if emp is None:
        conn.close()
        return False
    c.execute(
        "SELECT COUNT(*) FROM manual_overrides WHERE target_date = ? AND duty_group = ? AND emp_no = ?",
        (target_date, duty_group, emp_no))
    if c.fetchone()[0] == 0:
        c.execute(
            "INSERT INTO manual_overrides (target_date, duty_group, emp_no, reason) VALUES (?, ?, ?, ?)",
            (target_date, duty_group, emp_no, reason))
        conn.commit()
    conn.close()
    return True


def delete_manual_override(override_id: int) -> None:
    """수동 배정(벌당직/대근)을 취소한다."""
    conn = get_conn()
    c = conn.cursor()
    c.execute("DELETE FROM manual_overrides WHERE id = ?", (override_id,))
    conn.commit()
    conn.close()


def employee_change_date(value: str | None = None) -> str:
    """변동일을 검증하고, 생략 시 한국 시간의 오늘 날짜를 반환한다."""
    if not value:
        return datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=9))).date().isoformat()
    try:
        date = datetime.date.fromisoformat(value)
        if date.isoformat() != value:
            raise ValueError
    except ValueError as exc:
        raise ValueError('변동일은 YYYY-MM-DD 형식의 올바른 날짜로 입력하세요.') from exc
    return value


def _ensure_employee_change_history(c: sqlite3.Cursor) -> None:
    c.execute('''CREATE TABLE IF NOT EXISTS employee_change_history (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        emp_id INTEGER NOT NULL, emp_no TEXT NOT NULL, name TEXT NOT NULL,
        change_type TEXT NOT NULL, effective_date TEXT NOT NULL, recorded_at TEXT NOT NULL,
        reason TEXT NOT NULL, source TEXT NOT NULL,
        before_state TEXT, after_state TEXT NOT NULL
    )''')
    c.execute('CREATE INDEX IF NOT EXISTS idx_employee_changes_emp ON employee_change_history(emp_id, id)')


def _employee_snapshots(c: sqlite3.Cursor, emp_id: int | None = None) -> list[dict]:
    columns = ('id', 'emp_no', 'name', 'dept', 'duty_group', 'group_idx', 'is_active')
    query = 'SELECT ' + ','.join(columns) + ' FROM employees'
    rows = c.execute(query + (' WHERE id=?' if emp_id is not None else ''),
                     (emp_id,) if emp_id is not None else ()).fetchall()
    return [dict(zip(columns, row)) for row in rows]


def _record_employee_change(c: sqlite3.Cursor, before: dict | None, after: dict,
                            change_type: str, reason: str, effective_date: str,
                            source: str = '직원 관리') -> bool:
    # 실제 직원 순번은 유지하되 변동 이력의 전후 상태에서는 제외한다.
    before = {k: v for k, v in before.items() if k != 'group_idx'} if before is not None else None
    after = {k: v for k, v in after.items() if k != 'group_idx'}
    if before == after and change_type in ('정보변경', '엑셀변경'):
        return False
    _ensure_employee_change_history(c)
    recorded = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=9))).isoformat(timespec='seconds')
    c.execute('''INSERT INTO employee_change_history
                 (emp_id,emp_no,name,change_type,effective_date,recorded_at,reason,source,before_state,after_state)
                 VALUES (?,?,?,?,?,?,?,?,?,?)''',
              (after['id'], after['emp_no'], after['name'], change_type, effective_date, recorded,
               reason.strip() or '사유 미입력', source,
               json.dumps(before, ensure_ascii=False) if before is not None else None,
               json.dumps(after, ensure_ascii=False)))
    return True


def remove_employee_order_history() -> dict:
    """백업 후 순번만 바뀐 이력을 삭제하고 나머지 이력의 순번 필드를 제거한다."""
    conn = get_conn()
    try:
        stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S_%f')
        backup = os.path.join(os.path.dirname(os.path.abspath(DB_NAME)),
                              f'duty_system_v3_before_history_cleanup_{stamp}.db')
        with sqlite3.connect(backup) as dest:
            conn.backup(dest)
        deleted, updated = [], []
        with conn:
            conn.execute('BEGIN IMMEDIATE')
            rows = conn.execute('SELECT id,change_type,before_state,after_state FROM employee_change_history').fetchall()
            for row_id, kind, before_json, after_json in rows:
                before = json.loads(before_json) if before_json else None
                after = json.loads(after_json)
                clean_before = {k: v for k, v in before.items() if k != 'group_idx'} if before is not None else None
                clean_after = {k: v for k, v in after.items() if k != 'group_idx'}
                if (before is not None and kind in ('정보변경', '엑셀변경')
                        and before.get('group_idx') != after.get('group_idx')
                        and clean_before == clean_after):
                    conn.execute('DELETE FROM employee_change_history WHERE id=?', (row_id,))
                    deleted.append(row_id)
                elif before != clean_before or after != clean_after:
                    conn.execute('UPDATE employee_change_history SET before_state=?,after_state=? WHERE id=?',
                                 (json.dumps(clean_before, ensure_ascii=False) if clean_before is not None else None,
                                  json.dumps(clean_after, ensure_ascii=False), row_id))
                    updated.append(row_id)
        return {'backup': backup, 'deleted_ids': deleted, 'updated_ids': updated,
                'deleted': len(deleted), 'updated': len(updated), 'remaining': len(rows) - len(deleted)}
    finally:
        conn.close()


def register_employee(emp_no: str, name: str, dept: str, duty_group: str,
                      reason: str = '', effective_date: str | None = None) -> int:
    """전입자를 그룹의 마지막 순번에 등록한다. 기존 사번은 덮어쓰지 않는다."""
    emp_no, name, dept = emp_no.strip(), name.strip(), dept.strip()
    if len(emp_no) != 8 or not emp_no.isascii() or not emp_no.isdigit() or int(emp_no) == 0:
        raise ValueError('사번은 0으로만 구성되지 않은 8자리 숫자로 입력하세요.')
    if not name or name.startswith('*') or not dept:
        raise ValueError('성명과 부서를 입력하세요. 안내 문구는 직원으로 등록할 수 없습니다.')
    if duty_group not in EMPLOYEE_GROUPS:
        raise ValueError('올바른 당직그룹을 선택하세요.')
    change_date = employee_change_date(effective_date)
    conn = get_conn()
    try:
        with conn:
            # 최대 순번 조회와 삽입을 묶어 동시 등록 시 순번 중복을 방지한다.
            conn.execute('BEGIN IMMEDIATE')
            if conn.execute('SELECT 1 FROM employees WHERE emp_no = ?', (emp_no,)).fetchone():
                raise ValueError('이미 등록된 사번입니다. 전출자는 기존 직원의 상태를 활성으로 변경하세요.')
            index = conn.execute('SELECT COALESCE(MAX(group_idx), 0) + 1 FROM employees WHERE duty_group = ?',
                                 (duty_group,)).fetchone()[0]
            cursor = conn.execute(
                'INSERT INTO employees (emp_no,name,dept,duty_group,group_idx,is_active) VALUES (?,?,?,?,?,1)',
                (emp_no, name, dept, duty_group, index))
            emp_id = int(cursor.lastrowid)
            _record_employee_change(conn.cursor(), None, _employee_snapshots(conn.cursor(), emp_id)[0],
                                    '전입', reason, change_date)
            return emp_id
    finally:
        conn.close()


def transfer_out_employee(emp_id: int, reason: str = '', effective_date: str | None = None) -> bool:
    """전출자를 정규 배정에서 제외한다. 이력·순번·수동 배정은 보존한다."""
    change_date = employee_change_date(effective_date)
    conn = get_conn()
    try:
        with conn:
            conn.execute('BEGIN IMMEDIATE')
            c = conn.cursor()
            previous = _employee_snapshots(c, emp_id)
            if not previous:
                return False
            before = previous[0]
            c.execute("UPDATE source_roster_tracks SET is_active=0 WHERE emp_id=?", (emp_id,))
            _ensure_employee_change_history(c)
            last = c.execute('SELECT change_type FROM employee_change_history WHERE emp_id=? '
                             'ORDER BY id DESC LIMIT 1', (emp_id,)).fetchone()
            if before['is_active'] == 0 and last and last[0] == '전출':
                return True
            c.execute('UPDATE employees SET is_active = 0 WHERE id = ?', (emp_id,))
            _record_employee_change(c, before, _employee_snapshots(c, emp_id)[0], '전출', reason, change_date)
            return True
    finally:
        conn.close()


def update_employee(emp_id: int, duty_group: str, group_idx: int,
                    name: str, dept: str, is_active: int,
                    reason: str = '', effective_date: str | None = None) -> bool:
    """직원 정보를 수정한다(퇴사=제외, 승진=그룹/순번 변경 등). 대상 없으면 False."""
    change_date = employee_change_date(effective_date)
    conn = get_conn()
    try:
        with conn:
            conn.execute('BEGIN IMMEDIATE')
            c = conn.cursor()
            previous = _employee_snapshots(c, emp_id)
            if not previous:
                return False
            before = previous[0]
            c.execute('''UPDATE employees SET duty_group=?,group_idx=?,name=?,dept=?,is_active=? WHERE id=?''',
                      (duty_group, int(group_idx), str(name).strip(), dept, 1 if int(is_active) else 0, emp_id))
            after = _employee_snapshots(c, emp_id)[0]
            if before != after:
                if before['duty_group'] != after['duty_group'] or before['group_idx'] != after['group_idx']:
                    c.execute('DELETE FROM source_roster_tracks WHERE emp_id=?', (emp_id,))
                elif before['is_active'] != after['is_active']:
                    c.execute('UPDATE source_roster_tracks SET is_active=? WHERE emp_id=?',
                              (after['is_active'], emp_id))
                change_type = '정보변경'
                if before['is_active'] != after['is_active']:
                    change_type = '제외' if not after['is_active'] else '활성복구'
                    if after['is_active']:
                        _ensure_employee_change_history(c)
                        last = c.execute("SELECT change_type FROM employee_change_history WHERE emp_id=? "
                                         "AND change_type IN ('전출','제외','활성복구','재전입') ORDER BY id DESC LIMIT 1",
                                         (emp_id,)).fetchone()
                        if last and last[0] == '전출':
                            change_type = '재전입'
                _record_employee_change(c, before, after, change_type, reason, change_date)
            return True
    finally:
        conn.close()


def list_employee_changes(keyword: str = '', emp_id: int | None = None,
                          change_type: str = '', page: int = 1, page_size: int | None = 50) -> tuple[list[dict], int]:
    """기록 최신순 이력 조회. page_size=None이면 검색 결과 전체를 반환한다."""
    conn = get_conn()
    try:
        with conn:
            c = conn.cursor()
            _ensure_employee_change_history(c)
            conditions, params = [], []
            if keyword:
                conditions.append('(name LIKE ? OR emp_no LIKE ? OR before_state LIKE ?)')
                params.extend([f'%{keyword.strip()}%'] * 3)
            if emp_id is not None:
                conditions.append('emp_id=?')
                params.append(emp_id)
            if change_type:
                conditions.append('change_type=?')
                params.append(change_type)
            where = ' WHERE ' + ' AND '.join(conditions) if conditions else ''
            total = c.execute('SELECT COUNT(*) FROM employee_change_history' + where, params).fetchone()[0]
            query = 'SELECT * FROM employee_change_history' + where + ' ORDER BY id DESC'
            if page_size is None:
                c.execute(query, params)
            else:
                page = min(max(1, page), max(1, (total + page_size - 1) // page_size))
                c.execute(query + ' LIMIT ? OFFSET ?', [*params, page_size, (page - 1) * page_size])
            names = [column[0] for column in c.description]
            rows = [dict(zip(names, row)) for row in c.fetchall()]
            labels = {'name': '성명', 'emp_no': '사번', 'dept': '부서', 'duty_group': '당직그룹',
                      'is_active': '상태'}
            for row in rows:
                before = json.loads(row['before_state']) if row['before_state'] else {}
                after = json.loads(row['after_state'])
                changes = []
                for key, label in labels.items():
                    if before.get(key) != after.get(key):
                        old, new = before.get(key), after.get(key)
                        if key == 'is_active':
                            old = '활성' if old == 1 else '제외' if old == 0 else '—'
                            new = '활성' if new == 1 else '제외'
                        changes.append(f'{label}: {old if old is not None else "—"} → {new}')
                row['changes'] = changes
            return rows, total
    finally:
        conn.close()


def list_employees(duty_group: str | None = None, status: str = 'all',
                   keyword: str | None = None) -> list[tuple]:
    """직원 목록 조회. status: all|active|inactive, keyword: 사번/성명/부서 부분일치."""
    where: list[str] = []
    params: list[object] = []
    if duty_group and duty_group != 'all':
        where.append("duty_group = ?")
        params.append(duty_group)
    if status == 'active':
        where.append("is_active = 1")
    elif status == 'inactive':
        where.append("is_active = 0")
    if keyword:
        where.append("(name LIKE ? OR emp_no LIKE ? OR dept LIKE ?)")
        kw = f"%{keyword.strip()}%"
        params += [kw, kw, kw]
    sql = "SELECT id, duty_group, group_idx, name, emp_no, dept, is_active FROM employees"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY duty_group, group_idx"
    conn = get_conn()
    c = conn.cursor()
    c.execute(sql, params)
    rows = c.fetchall()
    conn.close()
    return rows


def get_duty_groups() -> list[str]:
    """현재 등록된 당직그룹 목록."""
    conn = get_conn()
    c = conn.cursor()
    c.execute("SELECT DISTINCT duty_group FROM employees ORDER BY duty_group")
    groups = [r[0] for r in c.fetchall()]
    conn.close()
    return groups


def get_assignments(year_month: str) -> list[tuple]:
    """월별 자동·수동 배정 현황. 날짜 → 일직/숙직 → 직급 → 이름순."""
    conn = get_conn()
    c = conn.cursor()
    # 수동 배정이 기록한 이력도 있으므로, 수동 배정(override)과 동일한 (날짜, 사번)의
    # 이력 행은 화면에서 제외하여 이중 표시를 방지한다.
    c.execute('''SELECT h.duty_date, h.duty_type, e.name, e.emp_no, e.duty_group, '배정' AS kind
                 FROM duty_history h JOIN employees e ON h.emp_id = e.id
                 WHERE h.duty_date LIKE ?
                   AND NOT EXISTS (SELECT 1 FROM manual_overrides m
                                   WHERE m.target_date = h.duty_date AND m.emp_no = e.emp_no)''',
              (year_month + '%',))
    rows = c.fetchall()
    c.execute('''SELECT m.target_date,
                        CASE WHEN INSTR(m.duty_group, '일직_') = 1 THEN '일직'
                             WHEN INSTR(m.duty_group, '숙직_') = 1 THEN '숙직'
                             ELSE '수동' END,
                        e.name, e.emp_no, m.duty_group, m.reason
                 FROM manual_overrides m JOIN employees e ON m.emp_no = e.emp_no
                 WHERE m.target_date LIKE ?''', (year_month + '%',))
    rows += c.fetchall()
    conn.close()
    # 수동 여부와 무관하게 일직을 먼저, 일직 안에서는 높은 직급부터 표시.
    # 2/3직급은 현재 명부의 통합 그룹이며 개별 직급 정보는 저장하지 않는다.
    type_order = {'일직': 0, '숙직': 1, '수동': 2}
    grade_order = {GROUP_FIRST: 1, GROUP_SEN: 2, GROUP_JUN: 4, GROUP_SUK: 3}
    return sorted(rows, key=lambda r: (r[0], type_order.get(r[1], 9),
                                       grade_order.get(r[4], 99), r[2], r[3]))


def export_assignments_excel(year_month: str) -> bytes:
    """화면과 같은 월별 배정을 원본 당직명령 형태의 XLSX로 출력한다(DB 쓰기 없음)."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins

    try:
        month = datetime.datetime.strptime(year_month, '%Y-%m').date()
    except ValueError as exc:
        raise ValueError('올바르지 않은 연월입니다 (YYYY-MM).') from exc
    if year_month != month.strftime('%Y-%m'):
        raise ValueError('올바르지 않은 연월입니다 (YYYY-MM).')
    rows = get_assignments(year_month)
    if not rows:
        raise ValueError(f'{year_month}에 다운로드할 배정이 없습니다.')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '당직명령'
    ws.sheet_view.showGridLines = False
    # docs 원본과 같은 B:I 열 구성, 열 너비, 제목·머리글·본문 높이.
    widths = {'A': 1.78, 'B': 6.55, 'C': 7.55, 'D': 6.33, 'E': 9.89,
              'F': 8.66, 'G': 10.89, 'H': 22.44, 'I': 13.22}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.merge_cells('B1:I1')
    ws['B1'] = f'{month.year}년 {month.month}월 당직명령'
    ws['B1'].font = Font(name='HY헤드라인M', size=20)
    ws['B1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32.25
    ws.row_dimensions[2].height = 9.95
    ws.row_dimensions[3].height = 20.25
    edge = Side(style='thin', color='000000')
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for col, title in enumerate(('일자', '요일', '구분', '당직비', '성명', '사번', '변경사항', '서명'), 2):
        cell = ws.cell(3, col, title)
        cell.font = Font(name='돋움', size=11, bold=True)
        cell.fill = PatternFill('solid', fgColor='EBF1DE')
        cell.border = border
        cell.alignment = center

    calendars = {}
    row_number = 4
    for duty_date, daily_rows in groupby(rows, key=lambda row: row[0]):
        daily = list(daily_rows)
        date = datetime.date.fromisoformat(duty_date[:10])
        if date.year not in calendars:
            calendars[date.year] = holidays.KR(years=[date.year])
        holiday = date.weekday() >= 5 or date in calendars[date.year]
        start = row_number
        for _, duty_type, name, emp_no, _, kind in daily:
            manual = kind != '배정'
            values = (date.day, '월화수목금토일'[date.weekday()], duty_type,
                      None, name, emp_no, kind if manual else None, None)
            for col, value in enumerate(values, 2):
                cell = ws.cell(row_number, col, value)
                # 이름·사유가 '='로 시작해도 수식으로 실행하지 않고 문자 그대로 내보낸다.
                if isinstance(value, str):
                    cell.data_type = 's'
                cell.font = Font(name='돋움', size=10 if col in (6, 7) else 11,
                                 bold=manual and col in (6, 7),
                                 color='FF0000' if col == 3 and holiday else '000000')
                cell.alignment = center
                cell.border = border
                if manual and col >= 6:
                    cell.fill = PatternFill('solid', fgColor='FFF2CC')
                if col == 7:
                    cell.number_format = '@'
            ws.row_dimensions[row_number].height = 15
            row_number += 1
        if row_number - start > 1:
            for col in ('B', 'C'):
                ws.merge_cells(f'{col}{start}:{col}{row_number - 1}')
        # 같은 날짜의 연속된 동일 구분을 원본처럼 병합한다(예: 일직 2명).
        group_start = start
        for _, same_type in groupby(daily, key=lambda row: row[1]):
            count = len(list(same_type))
            if count > 1:
                ws.merge_cells(start_row=group_start, start_column=4,
                               end_row=group_start + count - 1, end_column=4)
            group_start += count

    note_row = row_number + 1
    for offset, note in enumerate(('※ 당직비·서명은 확인 후 작성해 주세요.',
                                   '※ 굵은 글씨·노란색은 수동 배정이며, 변경사항에 사유를 표시합니다.')):
        ws.merge_cells(start_row=note_row + offset, start_column=2,
                       end_row=note_row + offset, end_column=9)
        cell = ws.cell(note_row + offset, 2, note)
        cell.font = Font(name='돋움', size=9)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[note_row + offset].height = 20
    ws.freeze_panes = 'F4'
    ws.print_area = f'B1:I{note_row + 1}'
    ws.print_title_rows = '1:3'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.24, right=0.24, top=0.25, bottom=0.25,
                                  header=0.1, footer=0.1)
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def regenerate_month(date_str: str) -> None:
    """지정 날짜가 속한 달 전체를 다시 배정한다(수동 배정 반영 등)."""
    d = datetime.datetime.strptime(date_str[:10], "%Y-%m-%d").date()
    last_day = calendar.monthrange(d.year, d.month)[1]
    generate_duty_plan(f"{d.year:04d}-{d.month:02d}-01", f"{d.year:04d}-{d.month:02d}-{last_day:02d}")


def list_months() -> list[tuple[str, int]]:
    """duty_history에 기록된 월별(YYYY-MM)과 기록 수. 월 삭제 화면 안내용."""
    conn = get_conn()
    c = conn.cursor()
    rows = c.execute('SELECT SUBSTR(duty_date,1,7), COUNT(*) FROM duty_history GROUP BY 1 ORDER BY 1').fetchall()
    conn.close()
    return [(str(r[0]), int(r[1])) for r in rows]


def _rewind_pointer(c: sqlite3.Cursor, pointer_key: str, is_weekend_slot: bool,
                    duty_type: str, month_start: str,
                    duty_group: str | None = None) -> None:
    """pointer_key를 '해당 웹 이전'의 마지막 배정자로 되돌린다. 이전 기록이 없으면 포인터 행 제거.

    duty_type은 '숙직' 또는 '일직'이며, is_weekend_slot으로 평일/휴무일 슬롯을 구분한다
    (숙직은 평일=ptr_sukjig_fwd, 휴무일=ptr_sukjig_rev — 공휴일 판정이 SQL로 불가하므로
    후보 행을 넉넉히 읽어 파이썬에서 판정한다).
    """
    grp = " AND e.duty_group = ?" if duty_group else ""
    sql = (f"SELECT h.duty_date, h.emp_id FROM duty_history h "
           f"JOIN employees e ON h.emp_id = e.id "
           f"WHERE h.duty_date < ? AND h.duty_type = ?{grp} "
           f"ORDER BY h.duty_date DESC, h.id DESC LIMIT 400")
    params: list[object] = [month_start, duty_type]
    if duty_group:
        params.append(duty_group)
    for duty_date, emp_id in c.execute(sql, params).fetchall():
        d = datetime.date.fromisoformat(str(duty_date)[:10])
        if (d.weekday() >= 5 or d in holidays.KR(years=[d.year])) == is_weekend_slot:
            c.execute("REPLACE INTO system_state (pointer_key, emp_id) VALUES (?, ?)", (pointer_key, emp_id))
            return
    c.execute("DELETE FROM system_state WHERE pointer_key = ?", (pointer_key,))


def delete_month(year_month: str) -> dict:
    """해당 월(YYYY-MM)의 생성 데이터를 지우고 '생성 전' 상태로 되돌린다.

    - duty_history: 해당 월 행 삭제
    - skip_queue: 해당 월 생성 중 만들어진 스킵 기록(skipped_date) 삭제.
      해당 월에 소비된 복귀자(is_returned=1, return_year_month=연월)는 다시 대기(0)로 복구
    - system_state 포인터:
      * 스냅샷(system_state_snapshot)이 있으면 생성 시작 시점 값으로 정확 복원
        → 지운 월을 다시 생성하면 처음과 동일한 결과(결정적).
      * 스냅샷이 없으면(이 기능 이전에 생성된 월) 이력 기반 추정 롤백 — 결과가 달라질 수 있음.
      * 해당 월 스냅샷은 소비 후 제거(재생성 시 새 스냅샷 기록). 이후 월 스냅샷도 상태가
        무효화되므로 함께 제거한다.
    - manual_overrides는 수동 입력 자산이므로 유지한다.
    반환: 처리 건수 요약 {'history_deleted', 'skip_cleared', 'returns_restored',
                          'pointer_restored_from': 'snapshot'|'history'|'none'}.
    """
    ym = year_month[:7]
    prefix = f"{ym}%"
    month_start = f"{ym}-01"
    conn = get_conn()
    init_db()
    c = conn.cursor()
    hist = int(c.execute("SELECT COUNT(*) FROM duty_history WHERE duty_date LIKE ?", (prefix,)).fetchone()[0])
    skips = int(c.execute("SELECT COUNT(*) FROM skip_queue WHERE skipped_date LIKE ?", (prefix,)).fetchone()[0])
    restored = int(c.execute(
        "SELECT COUNT(*) FROM skip_queue WHERE is_returned = 1 AND (return_year_month = ? OR returned_date LIKE ?)", (ym, ym + '%')).fetchone()[0])

    if hist:
        c.execute("SELECT snapshot FROM system_state_snapshot WHERE year_month = ?", (ym,))
        row = c.fetchone()
        if row and row[0]:
            c.execute("DELETE FROM system_state")
            for pair in str(row[0]).split(';'):
                if '=' in pair:
                    k, v = pair.split('=', 1)
                    c.execute("REPLACE INTO system_state (pointer_key, emp_id) VALUES (?, ?)", (k, int(v)))
            method: str = 'snapshot'
        else:
            _rewind_pointer(c, PTR_FWD, False, '숙직', month_start)
            _rewind_pointer(c, PTR_REV, True, '숙직', month_start)
            _rewind_pointer(c, PTR_SEN, True, '일직', month_start, GROUP_SEN)
            _rewind_pointer(c, PTR_JUN, True, '일직', month_start, GROUP_JUN)
            method = 'history'
    else:
        method = 'none'

    c.execute("UPDATE skip_queue SET is_returned = 0,returned_date=NULL WHERE is_returned = 1 AND (return_year_month = ? OR returned_date LIKE ?)", (ym, ym + '%'))
    c.execute("DELETE FROM skip_queue WHERE skipped_date LIKE ?", (prefix,))
    c.execute("DELETE FROM duty_history WHERE duty_date LIKE ?", (prefix,))
    c.execute("DELETE FROM system_state_snapshot WHERE year_month >= ?", (ym,))
    conn.commit()
    conn.close()
    print(f"[삭제] {ym} 이력 {hist}건, 스킵 {skips}건 제거, 복귀자 {restored}건 대기 복구 (포인터: {method})")
    return {'history_deleted': hist, 'skip_cleared': skips, 'returns_restored': restored,
            'pointer_restored_from': method}


def main() -> None:
    migrate_v2_to_v3()
    while True:
        print("\n=== 본사 당직명령 시스템 ===")
        print("1. 마스터 데이터 불러오기(엑셀)")
        print("2. 당직 배정 생성")
        print("3. 수동 예외 배정")
        print("4. 당직 현황 조회")
        print("0. 종료")
        menu: str = input("선택: ")

        if menu == '1':
            load_initial_data(MASTER_FILE)
        elif menu == '2':
            start: str = input("기간 시작 (YYYY-MM-DD): ")
            end: str = input("기간 종료 (YYYY-MM-DD): ")
            load_initial_data(MASTER_FILE)  # 최신 원부 반영 후 배정
            generate_duty_plan(start, end)
        elif menu == '3':
            add_manual_override()
        elif menu == '4':
            show_status()
        elif menu == '0':
            break


if __name__ == '__main__':
    main()
