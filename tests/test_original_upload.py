"""원본 엑셀 기준 DB 구성과 업로드의 보존/실패 동작 검증."""
import contextlib
import io
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

import dangjik_basic1 as dy
import app as web

SOURCE = Path(__file__).resolve().parents[1] / 'docs/2026년_7월_본사_당직명령.xlsx'


@unittest.skipUnless(SOURCE.exists(), '로컬 원본 엑셀이 필요합니다.')
class OriginalUploadTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.original_db = dy.DB_NAME
        dy.DB_NAME = str(Path(self.tmp.name) / 'test.db')
        dy.init_db()

    def tearDown(self) -> None:
        dy.DB_NAME = self.original_db
        self.tmp.cleanup()

    def dump(self) -> list[str]:
        with dy.get_conn() as c:
            return list(c.iterdump())

    def test_parse_identity_status_and_holiday_order(self) -> None:
        before = self.dump()
        r = dy.read_original_order(str(SOURCE))
        self.assertEqual((r['month'], len(r['employees']), len(r['duties'])), ('2026-07', 376, 49))
        target = next(d for d in r['duties'] if d['date'] == '2026-07-22')
        self.assertEqual(target['emp_no'], '20161566')
        rev = [x for x in r['tracks'] if x['pointer_key'] == dy.PTR_REV]
        self.assertEqual([x['emp_no'] for x in rev[:4]], ['12162197','13161933','17161202','17161431'])
        dispatch = [x for x in r['tracks'] if x['emp_no'] == '19161137']
        self.assertEqual({x['pointer_key']: x['is_active'] for x in dispatch}, {dy.PTR_FWD: 1, dy.PTR_REV: 0})
        self.assertEqual(self.dump(), before)

    def test_import_replaces_future_preserves_past_manual_settings_and_ids(self) -> None:
        keep = dy.register_employee('20161566', '이지수', '다른부서', dy.GROUP_SUK)
        outsider = dy.register_employee('99999999', '별도직원', '부서', dy.GROUP_JUN)
        dy.set_skip_current_month_only(False)
        dy.insert_manual_override('2026-08-01', dy.GROUP_JUN, '99999999', '대근')
        with dy.get_conn() as c:
            c.execute("INSERT INTO duty_history(duty_date,duty_type,emp_id) VALUES ('2026-06-01','숙직',?)", (keep,))
            c.execute("INSERT INTO duty_history(duty_date,duty_type,emp_id) VALUES ('2026-08-01','숙직',?)", (keep,))
            c.execute("INSERT INTO system_state_snapshot VALUES ('2026-08','stale')")
            history = c.execute('SELECT * FROM employee_change_history ORDER BY id').fetchall()
        result = dy.load_original_order(str(SOURCE))
        self.assertTrue(Path(result['backup']).exists())
        self.assertFalse(dy.get_skip_current_month_only())
        with dy.get_conn() as c:
            self.assertEqual(c.execute("SELECT id FROM employees WHERE emp_no='20161566'").fetchone()[0], keep)
            self.assertEqual(c.execute('SELECT is_active FROM employees WHERE id=?', (outsider,)).fetchone()[0], 0)
            self.assertEqual(c.execute('SELECT count(*) FROM duty_history').fetchone()[0], 50)
            self.assertEqual(c.execute("SELECT count(*) FROM duty_history WHERE duty_date LIKE '2026-08%'").fetchone()[0], 0)
            self.assertEqual(c.execute('SELECT count(*) FROM manual_overrides').fetchone()[0], 1)
            self.assertEqual(c.execute('SELECT count(*) FROM system_state_snapshot').fetchone()[0], 0)
            self.assertEqual(c.execute('SELECT * FROM employee_change_history ORDER BY id').fetchall()[:len(history)], history)
            self.assertEqual(c.execute("SELECT e.name FROM system_state s JOIN employees e ON e.id=s.emp_id WHERE s.pointer_key=?", (dy.PTR_FWD,)).fetchone()[0], '박상진')
            self.assertEqual(c.execute('PRAGMA integrity_check').fetchone()[0], 'ok')
        with contextlib.redirect_stdout(io.StringIO()):
            dy.regenerate_month('2026-08-01')
        with dy.get_conn() as c:
            self.assertEqual(c.execute("SELECT count(*) FROM duty_history WHERE duty_date LIKE '2026-08%'").fetchone()[0], 53)
        # 동일 원본 재등록 시 누적·중복 이력이나 직원변동 이력을 만들지 않는다.
        dy.load_original_order(str(SOURCE))
        with dy.get_conn() as c:
            change_count = c.execute('SELECT count(*) FROM employee_change_history').fetchone()[0]
        dy.load_original_order(str(SOURCE))
        with dy.get_conn() as c:
            self.assertEqual(c.execute('SELECT count(*) FROM duty_history').fetchone()[0], 50)
            self.assertEqual(c.execute('SELECT count(*) FROM employee_change_history').fetchone()[0], change_count)

    def test_validation_failure_does_not_modify_database(self) -> None:
        path = Path(self.tmp.name) / 'bad.xlsx'
        w = openpyxl.load_workbook(SOURCE)
        w['순서']['B4'] = w['순서']['B3'].value
        w.save(path); w.close()
        before = self.dump()
        with self.assertRaisesRegex(ValueError, '중복 사번'):
            dy.load_original_order(str(path))
        self.assertEqual(self.dump(), before)

    def test_write_failure_rolls_back_all_import_changes(self) -> None:
        before = self.dump()
        with patch.object(dy, '_record_employee_change', side_effect=RuntimeError('실패')):
            with self.assertRaises(RuntimeError):
                dy.load_original_order(str(SOURCE))
        self.assertEqual(self.dump(), before)

    def test_source_track_order_and_status_apply_and_manual_edits_release_them(self) -> None:
        dy.load_original_order(str(SOURCE))
        with dy.get_conn() as c:
            c.execute('DELETE FROM duty_history')
            number_to_id = dict(c.execute('SELECT emp_no,id FROM employees'))
            # 역순 두 번째는 평일 명부 역순의 이종준 대신 원본 휴일 명부의 강한국이다.
            c.execute('REPLACE INTO system_state VALUES (?,?)', (dy.PTR_REV, number_to_id['12162197']))
            chosen, _ = dy.get_next_candidate(c.cursor(), '2026-08-01', dy.GROUP_SUK, dy.PTR_REV, True)
            self.assertEqual(chosen, number_to_id['13161933'])
            # 조항성은 휴일에서는 제외되어 다음 직원으로 진행한다.
            # 실제 행 위치는 안내행 없이 매겨지므로 해당 직원 직전 순번을 조회한다.
            pos = c.execute('SELECT track_idx FROM source_roster_tracks WHERE pointer_key=? AND emp_id=?', (dy.PTR_REV,number_to_id['19161137'])).fetchone()[0]
            previous = c.execute('SELECT emp_id FROM source_roster_tracks WHERE pointer_key=? AND track_idx=?', (dy.PTR_REV,pos-1)).fetchone()[0]
            c.execute('REPLACE INTO system_state VALUES (?,?)', (dy.PTR_REV,previous))
            chosen, _ = dy.get_next_candidate(c.cursor(), '2026-08-01', dy.GROUP_SUK, dy.PTR_REV, True)
            self.assertNotIn(chosen, (number_to_id['19161137'],number_to_id['17161199']))
        dy.update_employee(number_to_id['19161137'],dy.GROUP_SUK,129,'조항성','부서',1)
        with dy.get_conn() as c:
            self.assertEqual(c.execute('SELECT count(*) FROM source_roster_tracks WHERE emp_id=?', (number_to_id['19161137'],)).fetchone()[0], 0)

    def test_status_change_preserves_original_order_and_new_arrival_enters_reverse(self) -> None:
        dy.load_original_order(str(SOURCE))
        with dy.get_conn() as c:
            emp=dy._employee_snapshots(c.cursor())
            target=next(e for e in emp if e['emp_no']=='19161137')
            before=c.execute('SELECT pointer_key,track_idx FROM source_roster_tracks WHERE emp_id=? ORDER BY pointer_key',(target['id'],)).fetchall()
        dy.update_employee(target['id'],target['duty_group'],target['group_idx'],target['name'],target['dept'],0)
        with dy.get_conn() as c:
            self.assertEqual(c.execute('SELECT pointer_key,track_idx FROM source_roster_tracks WHERE emp_id=? ORDER BY pointer_key',(target['id'],)).fetchall(),before)
            self.assertEqual(c.execute('SELECT sum(is_active) FROM source_roster_tracks WHERE emp_id=?',(target['id'],)).fetchone()[0],0)
        new_id=dy.register_employee('99999999','신규숙직','부서',dy.GROUP_SUK)
        with dy.get_conn() as c:
            c.execute('DELETE FROM system_state WHERE pointer_key=?',(dy.PTR_REV,))
            chosen,_=dy.get_next_candidate(c.cursor(),'2026-08-01',dy.GROUP_SUK,dy.PTR_REV,True)
            self.assertEqual(chosen,new_id)

    def test_web_original_upload_and_invalid_extension(self) -> None:
        client = web.app.test_client()
        with patch.object(web, 'UPLOAD_DIR', self.tmp.name):
            response = client.post('/upload', data={'file':(io.BytesIO(SOURCE.read_bytes()),'../../7월.xlsx')}, content_type='multipart/form-data',follow_redirects=True)
        body = response.get_data(as_text=True)
        self.assertIn('원본 2026-07 기준 등록 완료',body)
        self.assertIn('20161566',body)
        self.assertIn('평일·휴일',body)
        response = client.post('/upload', data={'file':(io.BytesIO(b'bad'),'old.xls')},content_type='multipart/form-data',follow_redirects=True)
        self.assertIn('.xlsx로 변환',response.get_data(as_text=True))

    def test_pending_source_skips_keep_return_month_and_track(self) -> None:
        path = SOURCE.with_name('2026년_8월_본사_당직명령.xlsx')
        if not path.exists():
            self.skipTest('8월 원본이 필요합니다.')
        result = dy.load_original_order(str(path))
        self.assertEqual(result['skips'], 34)
        with dy.get_conn() as c:
            self.assertEqual(c.execute("SELECT count(*) FROM skip_queue WHERE return_year_month='2026-09' AND is_returned=0").fetchone()[0], 23)
            self.assertEqual(c.execute("SELECT count(*) FROM skip_queue WHERE return_year_month='2026-10' AND is_returned=0").fetchone()[0], 11)
            self.assertEqual(c.execute("SELECT count(*) FROM skip_queue WHERE pointer_key=? AND return_year_month='2026-09'", (dy.PTR_FWD,)).fetchone()[0], 2)

    def test_legacy_template_dispatch_is_preserved(self) -> None:
        w=openpyxl.Workbook();w.active.title='직원명부(마스터)';w.create_sheet('초기당직이력(2월)')
        path=Path(self.tmp.name)/'legacy.xlsx';w.save(path);w.close()
        with patch.object(dy,'load_initial_data') as loader:
            result=dy.import_uploaded_workbook(str(path))
        loader.assert_called_once_with(str(path))
        self.assertEqual(result['format'],'template')


if __name__ == '__main__':
    unittest.main()
