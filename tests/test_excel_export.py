"""월별 엑셀 다운로드의 데이터 보존·수동 배정·서식·입력 오류 검증."""
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from urllib.parse import unquote

import openpyxl

import dangjik_basic1 as dy
from app import app


class ExcelExportTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.original_db = dy.DB_NAME
        dy.DB_NAME = str(Path(self.temp.name) / 'export.db')
        dy.init_db()
        self.senior = dy.register_employee('00123456', '일직정', '부서', dy.GROUP_SEN)
        self.junior = dy.register_employee('00234567', '일직부', '부서', dy.GROUP_JUN)
        self.night = dy.register_employee('00345678', '숙직자', '부서', dy.GROUP_SUK)
        with dy.get_conn() as conn:
            conn.executemany('INSERT INTO duty_history(duty_date,duty_type,emp_id) VALUES (?,?,?)', [
                ('2026-08-01', '일직', self.senior),
                ('2026-08-01', '일직', self.junior),
                ('2026-08-01', '숙직', self.night),
                ('2026-08-03', '숙직', self.night),
                ('2026-09-01', '숙직', self.night),
            ])
        dy.insert_manual_override('2026-08-03', dy.GROUP_SUK, '00345678', '대근')
        self.client = app.test_client()

    def tearDown(self) -> None:
        dy.DB_NAME = self.original_db
        self.temp.cleanup()

    def test_download_matches_screen_and_preserves_database(self) -> None:
        with dy.get_conn() as conn:
            before = list(conn.iterdump())
        response = self.client.get('/export/excel?month=2026-08')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype,
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('2026-08_본사_당직명령.xlsx', unquote(response.headers['Content-Disposition']))
        wb = openpyxl.load_workbook(BytesIO(response.data))
        ws = wb['당직명령']
        self.assertEqual(wb.sheetnames, ['당직명령'])
        self.assertEqual([ws.cell(r, 7).value for r in range(4, 8)],
                         [row[3] for row in dy.get_assignments('2026-08')])
        self.assertEqual([ws.cell(r, 6).value for r in range(4, 8)],
                         ['일직정', '일직부', '숙직자', '숙직자'])
        self.assertEqual(ws['G4'].data_type, 's')
        self.assertEqual(ws['G4'].value, '00123456')
        self.assertEqual(ws['H7'].value, '대근')
        self.assertTrue(ws['F7'].font.bold)
        self.assertIsNone(ws['E4'].value)  # 지급 기준이 저장되지 않은 당직비는 공란
        self.assertIsNone(ws['I4'].value)
        with dy.get_conn() as conn:
            self.assertEqual(list(conn.iterdump()), before)
        self.assertIn('formaction="/export/excel"',
                      self.client.get('/?month=2026-08').get_data(as_text=True))
        wb.close()

    def test_original_layout_merges_and_full_month_print_area(self) -> None:
        wb = openpyxl.load_workbook(BytesIO(dy.export_assignments_excel('2026-08')))
        ws = wb.active
        self.assertEqual(ws['B1'].value, '2026년 8월 당직명령')
        self.assertEqual([ws.cell(3, col).value for col in range(2, 10)],
                         ['일자', '요일', '구분', '당직비', '성명', '사번', '변경사항', '서명'])
        for area in ('B1:I1', 'B4:B6', 'C4:C6', 'D4:D5'):
            self.assertIn(area, ws.merged_cells)
        self.assertEqual(ws['B7'].value, 3)
        self.assertEqual(ws['C4'].font.color.rgb, '00FF0000')
        self.assertEqual(ws['C7'].font.color.rgb, '00000000')
        self.assertEqual(ws['F4'].border.bottom.style, 'thin')
        self.assertIn('$B$1:$I$10', ws.print_area)
        self.assertEqual(str(ws.page_setup.paperSize), ws.PAPERSIZE_A4)
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        wb.close()

    def test_invalid_or_empty_month_returns_message(self) -> None:
        for month in ('', 'bad', '2026-13', '2026-8', '../../etc/passwd'):
            with self.subTest(month=month):
                self.assertEqual(self.client.get('/export/excel', query_string={'month': month}).status_code, 302)
        response = self.client.get('/export/excel?month=2026-10', follow_redirects=True)
        self.assertIn('다운로드할 배정이 없습니다', response.get_data(as_text=True))

    def test_user_text_is_exported_as_literal_text(self) -> None:
        with dy.get_conn() as conn:
            conn.execute('UPDATE employees SET name=? WHERE id=?', ('=1+1', self.night))
            conn.execute('UPDATE manual_overrides SET reason=?', ('=2+2',))
        wb = openpyxl.load_workbook(BytesIO(dy.export_assignments_excel('2026-08')))
        for address, expected in (('F7', '=1+1'), ('H7', '=2+2')):
            self.assertEqual(wb.active[address].value, expected)
            self.assertEqual(wb.active[address].data_type, 's')
        wb.close()
