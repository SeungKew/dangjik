"""회색/초록/하늘색을 제외로 오해하지 않도록 원본 변환 규칙을 검증한다."""

from pathlib import Path
import tempfile
import unittest

import openpyxl
from openpyxl.styles import PatternFill

import dangjik_basic1 as dy
from scripts.correct_february_status import correct_template, database_changes, source_statuses


class StatusCorrectionTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.original_db = dy.DB_NAME
        dy.DB_NAME = str(Path(self.tmp.name) / 'test.db')
        dy.init_db()

    def tearDown(self) -> None:
        dy.DB_NAME = self.original_db
        self.tmp.cleanup()

    def test_only_pink_means_excluded(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '순서'
        colors = ['FFFF99CC', 'FFCCFFFF', 'FFCCFFCC', 'FFFFFFFF', None]
        for i, color in enumerate(colors, 3):
            ws.cell(i, 1, f'직원{i}')
            ws.cell(i, 2, f'{i:08d}')
            if color:
                ws.cell(i, 1).fill = PatternFill('solid', fgColor=color)
        path = Path(self.tmp.name) / 'source.xlsx'
        wb.save(path)
        statuses = source_statuses(path)
        self.assertEqual([s['active'] for s in statuses.values()], [False, True, True, True, True])
        ws.cell(8, 1, '중복 직원')
        ws.cell(8, 2, '00000003')
        wb.save(path)
        with self.assertRaises(ValueError):
            source_statuses(path)
        wb.close()

    def test_template_only_changes_status_and_reason_and_is_idempotent(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '직원명부(마스터)'
        ws.append(['당직그룹', '순번', '성명', '사번', '부서', '상태', '제외사유'])
        ws.append([dy.GROUP_SUK, 1, '직원', '00000001', '부서', '제외', '색상표시됨(제외)'])
        ws.append([dy.GROUP_SUK, 2, '* 안내', '00000000', None, '활성', None])
        ws2 = wb.create_sheet('초기당직이력(2월)')
        ws2.append(['2026-02-01', '숙직', '00000001'])
        before_history = list(ws2.values)
        before_identity = [tuple(c.value for c in r[:5]) for r in ws]
        source = {'00000001': dict(active=True, group=dy.GROUP_SUK, names=['직원'])}
        self.assertEqual(len(correct_template(wb, source)), 2)
        self.assertEqual(ws['F2'].value, '활성')
        self.assertIsNone(ws['G2'].value)
        self.assertEqual(ws['F3'].value, '제외')
        self.assertEqual(before_identity, [tuple(c.value for c in r[:5]) for r in ws])
        self.assertEqual(before_history, list(ws2.values))
        self.assertEqual(correct_template(wb, source), [])
        wb.close()

    def test_database_plan_preserves_new_arrivals_and_identity(self) -> None:
        emp = dy.register_employee('00000001', '직원', '부서', dy.GROUP_SUK)
        dy.transfer_out_employee(emp)
        dy.register_employee('00000002', '신규전입', '부서', dy.GROUP_SUK)
        source = {'00000001': dict(active=True, group=dy.GROUP_SUK, names=['직원'])}
        with dy.get_conn() as conn:
            changes = database_changes(conn, source)
            self.assertEqual(len(changes), 1)
            self.assertEqual(changes[0]['id'], emp)
            self.assertEqual(changes[0]['new'], 1)
            self.assertEqual(conn.execute('SELECT is_active FROM employees WHERE id=?', (emp,)).fetchone()[0], 0)
            source['00000001']['names'] = ['다른사람']
            with self.assertRaises(ValueError):
                database_changes(conn, source)


if __name__ == '__main__':
    unittest.main()
